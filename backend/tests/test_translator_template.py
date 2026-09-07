"""下载模板并在隔离服务上验证填写、重传和更新。"""
import ast
import io
import os
import time
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

import test_acceptance as api


def fill_row(sheet, values, row=2):
    columns = {
        str(cell.value).replace("*", ""): cell.column
        for cell in sheet[1] if cell.value
    }
    for label, value in values.items():
        sheet.cell(row, columns[label], value)


def upload(workbook, token):
    buffer = io.BytesIO()
    workbook.save(buffer)
    raw, content_type = api.mp(buffer.getvalue())
    return api.req("POST", "/api/import/translators", raw=raw, ct=content_type, token=token)[1]


def main():
    if os.getenv("BASE"):
        raise RuntimeError("模板专项只允许临时数据库；请移除 BASE。")
    api.start_isolated_server()
    try:
        for _ in range(40):
            try:
                api.req("GET", "/api/overview")
                break
            except OSError:
                time.sleep(0.2)
        else:
            raise RuntimeError("隔离服务启动失败")
        editor = api.req("POST", "/api/login", {"user": "资源端"})[1]["token"]
        viewer = api.req("POST", "/api/login", {"user": "boss"})[1]["token"]
        status, template_bytes = api.req("GET", "/api/export/translator-template")
        assert status == 200
        asset = Path(__file__).resolve().parents[1] / "app/templates/translator_import_template.xlsx"
        assert template_bytes == asset.read_bytes()
        book = load_workbook(io.BytesIO(template_bytes))
        assert book.sheetnames == ["译员导入", "项目经历", "名称映射", "填写说明"]
        parser = ast.parse((asset.parent.parent / "routers/admin.py").read_text())
        fields = next(
            ast.literal_eval(node.value) for node in parser.body
            if isinstance(node, ast.Assign)
            and any(isinstance(target, ast.Name) and target.id == "IMPORT_FIELDS" for target in node.targets)
        )
        expected_headers = ["译员ID"] + [label + ("*" if required else "") for label, _, required, _ in fields]
        assert [cell.value for cell in book["译员导入"][1]] == expected_headers
        for title in book.sheetnames[:3]:
            sheet = book[title]
            assert all(value is None for row in sheet.iter_rows(min_row=2, values_only=True) for value in row)
            assert sheet.freeze_panes == "C2", (title, sheet.freeze_panes)
        assert len(book["译员导入"].data_validations.dataValidation) == 12
        assert book["译员导入"]["D2"].number_format == "yyyy-mm-dd"
        print("PASS 下载文件与打包资源一致，表头、空白数据区、日期、下拉和冻结窗格正确")

        empty = upload(book, editor)
        assert empty["imported"] == 0 and empty["updated"] == 0 and not empty["invalid_rows"]
        print("PASS 空白模板不新增或覆盖译员")

        fill_row(book["译员导入"], {
            "姓名": "模板测试译员", "母语": "中文", "入库日期": date(2026, 9, 1),
            "性别": "女", "邮箱": "template-test@example.com", "主体类型": "个人译员",
            "语言对": "ZH→EN,ZH→JA", "日产字数": 4500,
            "人工评级": "A", "人工评级原因": "模板专项测试",
        })
        fill_row(book["项目经历"], {
            "译员邮箱": "template-test@example.com", "合作来源": "外部合作",
            "项目状态": "过往", "项目名称": "模板测试项目", "源语言": "ZH", "目标语言": "EN",
        })
        fill_row(book["名称映射"], {"译员邮箱": "template-test@example.com", "名称映射": "模板测试别名"})
        book.active = book.sheetnames.index("填写说明")
        result = upload(book, editor)
        assert result["imported"] == result["imported_projects"] == result["imported_aliases"] == 1, result
        assert not any(result[key] for key in ("invalid_rows", "invalid_project_rows", "invalid_alias_rows")), result
        profile = next(row for row in api.req("GET", "/api/translators")[1] if row["email"] == "template-test@example.com")
        assert profile["gender"] == "female" and profile["entity_type"] == "individual"
        assert profile["status"] == "Active" and profile["settlement_mode"] == "monthly"
        assert profile["daily_output"] == 4500 and profile["onboarding_date"] == "2026-09-01"
        assert profile["manual_rating"] == "A" and profile["manual_rating_reason"] == "模板专项测试"
        pairs = api.req("GET", f"/api/translators/{profile['id']}/language-pairs")[1]
        assert {(p["source_lang"], p["target_lang"]) for p in pairs} == {("ZH", "EN"), ("ZH", "JA")}
        print("PASS 保存于说明页仍正确导入主表、项目和名称映射，数值、日期和默认值正确")

        repeated = upload(book, editor)
        assert repeated["imported"] == 0 and repeated["skipped_dup_email"] == 1, repeated
        assert repeated["skipped_dup_projects"] == repeated["skipped_dup_aliases"] == 1, repeated
        print("PASS 重复上传跳过重复邮箱、项目和名称映射")

        assert api.code(lambda: upload(book, viewer)) == 403
        assert api.code(lambda: upload(book, None)) == 403
        print("PASS 只读及未登录用户不能导入")

        bad = load_workbook(io.BytesIO(template_bytes))
        fill_row(bad["译员导入"], {"姓名": "未填写性别", "母语": "中文", "入库日期": date(2026, 9, 1)})
        invalid = upload(bad, editor)
        assert invalid["imported"] == 0 and len(invalid["invalid_rows"]) == 1, invalid
        assert "gender" in invalid["invalid_rows"][0]["error"]
        print("PASS 缺少性别时拒绝该行并返回原因")

        updated = load_workbook(io.BytesIO(template_bytes))
        fill_row(updated["译员导入"], {
            "译员ID": profile["id"], "姓名": profile["name"], "母语": "中文",
            "入库日期": date(2026, 9, 1), "性别": "女", "邮箱": profile["email"],
            "状态": "Active", "结算策略": "累计结", "所在地": "上海",
        })
        update_result = upload(updated, editor)
        assert update_result["updated"] == 1 and update_result["imported"] == 0, update_result
        after = next(row for row in api.req("GET", "/api/translators")[1] if row["id"] == profile["id"])
        assert after["location"] == "上海" and after["settlement_mode"] == "cumulative"
        print("PASS 同ID更新保持编号，中文结算策略正确转换")
        print("7/7 template workflow checks passed")
        return 0
    finally:
        api.stop_isolated_server()


if __name__ == "__main__":
    raise SystemExit(main())
