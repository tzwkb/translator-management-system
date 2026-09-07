"""译员管理系统 demo 验收测试（可重跑）。

直接运行：backend/.venv/bin/python backend/tests/test_acceptance.py
默认会启动临时 SQLite + 临时端口的隔离后端；若设置 BASE，则改为打指定外部服务。
覆盖：鉴权/RBAC、边界健壮、PRD 硬规则（邮箱唯一/金额验算/汇总/汇总字段自动算）、
安全（加密脱敏/明文权限）、agent 护栏（钱相关进待审）。绕本机 Clash 代理直连。
"""
import atexit
import io
import json
import os
import socket
import subprocess
import sys
import tempfile
import time
from calendar import monthrange
from datetime import date, timedelta
from pathlib import Path
from urllib.parse import urlencode
import urllib.error
import urllib.request

from openpyxl import Workbook, load_workbook

BASE = os.getenv("BASE")
OP = urllib.request.build_opener(urllib.request.ProxyHandler({}))
SERVER = None
TMPDIR = None


def req(m, p, body=None, token=None, raw=None, ct="application/json", headers=None):
    h = dict(headers or {})
    if token:
        h["Authorization"] = "Bearer " + token
    if raw is not None:
        data, h["Content-Type"] = raw, ct
    elif body is not None:
        data, h["Content-Type"] = json.dumps(body).encode(), ct
    else:
        data = None
    r = urllib.request.Request(BASE + p, data=data, method=m, headers=h)
    with OP.open(r, timeout=8) as x:
        c = x.headers.get("content-type", "")
        return x.status, (json.loads(x.read().decode() or "null") if "json" in c else x.read())


def code(fn):
    try:
        fn()
        return "no-error"
    except urllib.error.HTTPError as e:
        return e.code


def mk_xlsx(rows):
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def mk_xlsx_sheets(sheets):
    wb = Workbook()
    for index, (title, rows) in enumerate(sheets.items()):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        for row in rows:
            ws.append(row)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def mp(xb):
    bd = "--x"
    head = (f"--{bd}\r\nContent-Disposition: form-data; name=\"file\"; "
            f"filename=\"x.xlsx\"\r\nContent-Type: application/octet-stream\r\n\r\n").encode()
    return head + xb + f"\r\n--{bd}--\r\n".encode(), f"multipart/form-data; boundary={bd}"


def mp_file(data, filename, content_type, fields=None):
    boundary = "acceptance-boundary"
    parts = []
    for key, value in (fields or {}).items():
        parts.append(
            f"--{boundary}\r\nContent-Disposition: form-data; name=\"{key}\"\r\n\r\n"
            f"{value}\r\n".encode()
        )
    parts.append(
        f"--{boundary}\r\nContent-Disposition: form-data; name=\"file\"; "
        f"filename=\"{filename}\"\r\nContent-Type: {content_type}\r\n\r\n".encode()
        + data + b"\r\n"
    )
    parts.append(f"--{boundary}--\r\n".encode())
    return b"".join(parts), f"multipart/form-data; boundary={boundary}"


def free_port():
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


def stop_isolated_server():
    global SERVER, TMPDIR
    if SERVER and SERVER.poll() is None:
        SERVER.terminate()
        try:
            SERVER.wait(timeout=5)
        except subprocess.TimeoutExpired:
            SERVER.kill()
            SERVER.wait(timeout=5)
    if TMPDIR:
        TMPDIR.cleanup()
    SERVER = None
    TMPDIR = None


def start_isolated_server():
    global BASE, SERVER, TMPDIR
    if BASE:
        return
    backend_dir = Path(__file__).resolve().parents[1]
    TMPDIR = tempfile.TemporaryDirectory(prefix="translator-acceptance-")
    db_path = Path(TMPDIR.name) / "acceptance.db"
    port = free_port()
    BASE = f"http://127.0.0.1:{port}"
    env = os.environ.copy()
    env["DB_URL"] = f"sqlite:///{db_path}"
    env["JWT_SECRET"] = "acceptance-test-secret"
    env["AES_KEY"] = "0" * 64
    env["TOKEN_TTL"] = "3600"
    env["SEED_DEMO_DATA"] = "1"
    env["UPLOAD_DIR"] = str(Path(TMPDIR.name) / "uploads")
    migration = subprocess.run(
        [sys.executable, "-m", "alembic", "-c", "alembic.ini", "upgrade", "head"],
        cwd=backend_dir,
        env=env,
        capture_output=True,
        text=True,
    )
    if migration.returncode:
        raise RuntimeError(migration.stdout + migration.stderr)
    SERVER = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "app.main:app", "--host", "127.0.0.1", "--port", str(port)],
        cwd=backend_dir,
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
    )
    atexit.register(stop_isolated_server)


def server_output():
    if not SERVER or not SERVER.stdout:
        return ""
    if SERVER.poll() is None:
        return "服务进程仍在运行，但未在超时时间内响应 /api/overview"
    try:
        return SERVER.stdout.read() or ""
    except Exception:
        return ""


def main():
    start_isolated_server()
    for _ in range(40):
        try:
            OP.open(BASE + "/api/overview", timeout=2)
            break
        except Exception:
            if SERVER and SERVER.poll() is not None:
                print("✗ 隔离服务启动失败")
                print(server_output())
                return 1
            time.sleep(0.3)
    else:
        print("✗ 服务没起来")
        print(server_output())
        return 1

    ok = []

    def chk(n, c, e=""):
        ok.append(bool(c))
        print(("✓" if c else "✗ 失败:"), n, e)

    _, ED = req("POST", "/api/login", {"user": "资源端"}); ET = ED["token"]
    _, BO = req("POST", "/api/login", {"user": "boss"}); BT = BO["token"]
    _, AG = req("POST", "/api/login", {"user": "资源端Agent"}); AT = AG["token"]
    run_tag = str(int(time.time() * 1000) % 100000)

    print("=== A 鉴权/RBAC ===")
    chk("登录不存在用户404", code(lambda: req("POST", "/api/login", {"user": "鬼"})) == 404)
    chk("坏token写403", code(lambda: req("POST", "/api/translators", {"name": "X"}, token="bad")) == 403)
    chk("无token写403", code(lambda: req("POST", "/api/translators", {"name": "X"})) == 403)
    chk("boss写403", code(lambda: req(
        "PUT",
        "/api/translators/1/capacity/override?month=2026-08",
        {"status": "健康", "reason": "权限测试"},
        token=BT,
    )) == 403)
    chk("agent不能批准403", code(lambda: req("POST", "/api/pending/1/approve", token=AT)) == 403)
    chk("viewer看审计403", code(lambda: req("GET", "/api/audit", token=BT)) == 403)

    print("=== B 边界/健壮 ===")
    chk("缺name 422", code(lambda: req("POST", "/api/translators", {"email": "a@b.c"}, token=ET)) == 422)
    chk("非法PO状态400", code(lambda: req("PUT", "/api/po/1/status", {"status": "乱填"}, token=ET)) == 400)
    chk("改不存在PO状态404", code(lambda: req("PUT", "/api/po/99999/status", {"status": "已支付"}, token=ET)) == 404)
    chk("操作不存在译员404", code(lambda: req("POST", "/api/translators/9999/rate-changes", {"change_date": "2026-06-01"}, token=ET)) == 404)
    chk("批准不存在待审404", code(lambda: req("POST", "/api/pending/9999/approve", token=ET)) == 404)
    chk("不存在译员子表GET空", req("GET", "/api/translators/9999/quality")[1] == [])
    chk("空月PO汇总不崩", isinstance(req("GET", "/api/po/summary?month=1999-01")[1]["by_currency"], dict))
    bad_raw, bad_ct = mp(mk_xlsx([["乱"]]))
    chk("LQE非汇总表400", code(lambda: req("POST", "/api/import/lqe", raw=bad_raw, token=ET, ct=bad_ct)) == 400)

    print("=== C PRD 硬规则 ===")
    dup_email = f"dup{run_tag}@x.com"
    new_email = f"new{run_tag}@x.com"
    req("POST", "/api/translators", {"name": "唯A", "email": dup_email, "native_language": "中文", "onboarding_date": "2025-01-01", "gender": "female"}, token=ET)
    chk("create重复邮箱400", code(lambda: req("POST", "/api/translators", {"name": "唯B", "email": dup_email, "native_language": "中文", "onboarding_date": "2025-01-01", "gender": "female"}, token=ET)) == 400)
    chk("update改成已存在邮箱400", code(lambda: req("PUT", "/api/translators/1", {"name": "张明", "email": dup_email, "native_language": "中文", "onboarding_date": "2025-01-01", "gender": "male"}, token=ET)) == 400)
    imp_raw, imp_ct = mp(mk_xlsx([
        ["name", "email", "native_language", "onboarding_date", "性别", "主体类型", "内部评级"],
        ["导入新", new_email, "中文", "2025-01-01", "女", "个人译员", "S"],
        ["导入重", dup_email, "中文", "2025-01-01", "女"],
        ["导入坏", f"bad{run_tag}@x.com", "中文", "2025-01-32", "女"],
    ]))
    _, imp = req("POST", "/api/import/translators", raw=imp_raw, token=ET, ct=imp_ct)
    chk("导入按邮箱去重(2行只进1)", imp["imported"] == 1, imp)
    chk("导入非法日期跳过并报告", len(imp.get("invalid_rows", [])) == 1, imp)
    imported_profile = next(
        item for item in req("GET", "/api/translators")[1]
        if item.get("email") == new_email
    )
    chk("导入支持中文性别和主体类型",
        imported_profile.get("gender") == "female"
        and imported_profile.get("entity_type") == "individual",
        imported_profile)
    chk("Excel 手工评级列不导入",
        imported_profile.get("internal_rating") is None,
        imported_profile)
    translator_count_before_update = len(req("GET", "/api/translators")[1])
    same_id_raw, same_id_ct = mp(mk_xlsx([
        [
            "译员ID", "姓名", "邮箱", "母语", "入库日期",
            "性别", "主体类型", "所在地",
        ],
        [
            imported_profile["id"], "导入新（已改名）", new_email, "中文",
            "2025-01-01", "女", "供应商", "上海",
        ],
    ]))
    _, same_id_result = req(
        "POST", "/api/import/translators",
        raw=same_id_raw, token=ET, ct=same_id_ct,
    )
    translators_after_update = req("GET", "/api/translators")[1]
    same_id_profile = next(
        item for item in translators_after_update
        if item["id"] == imported_profile["id"]
    )
    same_id_aliases = req(
        "GET",
        f"/api/translators/{imported_profile['id']}/aliases",
    )[1]
    chk(
        "译员Excel同ID覆盖且不新增记录",
        same_id_result.get("updated") == 1
        and same_id_result.get("imported") == 0
        and len(translators_after_update) == translator_count_before_update
        and same_id_profile["name"] == "导入新（已改名）"
        and same_id_profile["entity_type"] == "vendor"
        and same_id_profile["location"] == "上海",
        {"result": same_id_result, "profile": same_id_profile},
    )
    chk(
        "同ID改名自动保留旧姓名映射",
        any(item["alias"] == "导入新" for item in same_id_aliases),
        same_id_aliases,
    )
    duplicate_id_raw, duplicate_id_ct = mp(mk_xlsx([
        ["译员ID", "姓名", "邮箱", "母语", "入库日期", "性别"],
        [
            imported_profile["id"], "同ID第一行", new_email,
            "中文", "2025-01-01", "女",
        ],
        [
            imported_profile["id"], "同ID第二行", new_email,
            "中文", "2025-01-01", "女",
        ],
    ]))
    _, duplicate_id_result = req(
        "POST", "/api/import/translators",
        raw=duplicate_id_raw, token=ET, ct=duplicate_id_ct,
    )
    chk(
        "同一Excel内重复译员ID报告错误",
        duplicate_id_result.get("updated") == 1
        and any(
            "同一文件译员ID重复" in row.get("error", "")
            for row in duplicate_id_result.get("invalid_rows", [])
        ),
        duplicate_id_result,
    )
    _, po = req("GET", "/api/po?month=2026-06")
    chk("金额验算抓1行(410≠400)", len([p for p in po if not p["amount_ok"]]) == 1)
    _, sm = req("GET", "/api/po/summary?month=2026-06")
    chk("待付=Σ未开票+Σ已开票待付=2500", sm["by_currency"]["CNY"]["unpaid"] == 2500)
    req("POST", "/api/translators/2/rate-changes", {"change_date": "2026-06-25", "task_type": "翻译", "new_rate": 260}, token=ET)
    chk("翻译调价同步主表260", [t for t in req("GET", "/api/translators")[1] if t["id"] == 2][0]["translation_rate"] == 260)
    lp_target = "KO"
    req("POST", "/api/translators/1/language-pairs",
        {"source_lang": "ZH", "target_lang": lp_target, "translation_rate": 180, "review_rate": 90, "currency": "CNY"},
        token=ET)
    _, lps = req("GET", "/api/translators/1/language-pairs")
    zhko = next((x for x in lps if x["source_lang"] == "ZH" and x["target_lang"] == lp_target), {})
    chk("语言对可保存独立翻译费率", zhko.get("translation_rate") == 180, zhko)
    req("POST", "/api/translators/1/rate-changes",
        {"change_date": "2026-06-26", "source_lang": "ZH", "target_lang": lp_target, "task_type": "翻译", "new_rate": 188},
        token=ET)
    _, lps = req("GET", "/api/translators/1/language-pairs")
    zhko = next((x for x in lps if x["source_lang"] == "ZH" and x["target_lang"] == lp_target), {})
    chk("按语言对调价只更新该语言对", zhko.get("translation_rate") == 188, zhko)
    _, lp_po = req("POST", "/api/po",
                   {"translator_id": 1, "settlement_month": "2026-08", "source_lang": "ZH", "target_lang": lp_target,
                    "role": "翻译", "word_count": 1000, "po_number": "PO-LP-" + run_tag},
                   token=ET)
    chk("PO按语言对自动取价", lp_po.get("rate") == 188 and lp_po.get("amount") == 188 and lp_po.get("source_lang") == "ZH", lp_po)
    po_imp_raw, po_imp_ct = mp(mk_xlsx([
        ["译员", "结算月", "项目", "源语言", "目标语言", "角色", "字数", "单价", "币种", "状态", "PO号"],
        ["张明", "2026-09", "导入PO", "ZH", "EN", "翻译", 2000, None, "CNY", "未开票", "PO-IMP-" + run_tag],
        ["张明", "2026-09", "重复PO", "ZH", "EN", "翻译", 1000, 180, "CNY", "未开票", "PO-IMP-" + run_tag],
        ["不存在", "2026-09", "坏行", "ZH", "EN", "翻译", 1000, 180, "CNY", "未开票", "PO-BAD-" + run_tag],
    ]))
    try:
        _, po_imp = req("POST", "/api/import/po", raw=po_imp_raw, token=ET, ct=po_imp_ct)
    except urllib.error.HTTPError as e:
        po_imp = {"error": e.code}
    chk("PO导入支持有效/重复/错误行报告",
        po_imp.get("imported") == 1 and po_imp.get("skipped_dup_po") == 1 and len(po_imp.get("invalid_rows", [])) == 1,
        po_imp)
    _, imp_po_rows = req("GET", "/api/po?month=2026-09")
    imported_po = next((p for p in imp_po_rows if p.get("po_number") == "PO-IMP-" + run_tag), {})
    chk("PO导入可按语言对自动取价",
        imported_po.get("rate") == 180 and imported_po.get("amount") == 360,
        imported_po)
    chk("非法语言对400", code(lambda: req("POST", "/api/translators/1/language-pairs",
        {"source_lang": "ZH", "target_lang": "XX", "translation_rate": 1}, token=ET)) == 400)
    _, free_lp = req("POST", "/api/translators/1/language-pairs",
        {"source_lang": "EN", "target_lang": "JA", "translation_rate": 222, "currency": "CNY"},
        token=ET)
    chk("固定语言选项可自由组合", free_lp.get("source_lang") == "EN" and free_lp.get("target_lang") == "JA", free_lp)
    chk("源目标相同语言对400", code(lambda: req("POST", "/api/translators/1/language-pairs",
        {"source_lang": "ZH", "target_lang": "ZH", "translation_rate": 1}, token=ET)) == 400)

    print("=== C2 项目经历与基础字段 ===")
    profile_body = {
        "name": "项目经历校验" + run_tag,
        "email": f"experience{run_tag}@x.com",
        "native_language": "中文",
        "onboarding_date": "2026-07-01",
        "gender": "male",
        "entity_type": "individual",
        "current_project": "旧字段项目",
        "role": "翻译",
    }
    _, profile = req("POST", "/api/translators", profile_body, token=ET)
    experience_tid = profile["id"]
    chk("性别和主体类型可新增读取",
        profile.get("gender") == "male" and profile.get("entity_type") == "individual",
        profile)
    updated_profile_body = dict(
        profile_body, gender="female", entity_type="vendor",
    )
    _, updated_profile = req(
        "PUT", f"/api/translators/{experience_tid}", updated_profile_body, token=ET,
    )
    chk("性别和主体类型可编辑读取",
        updated_profile.get("gender") == "female"
        and updated_profile.get("entity_type") == "vendor",
        updated_profile)

    _, current_our = req(
        "POST", f"/api/translators/{experience_tid}/project-experiences",
        {
            "cooperation_source": "our_company",
            "project_status": "current",
            "project_name": " 我司当前项目 ",
            "role": "翻译",
            "source_lang": "zh",
            "target_lang": "en",
            "start_date": "2026-07-01",
            "remaining_volume": 12000,
            "deadline": "2026-08-01",
        },
        token=ET,
    )
    _, current_external = req(
        "POST", f"/api/translators/{experience_tid}/project-experiences",
        {
            "cooperation_source": "external",
            "project_status": "current",
            "project_name": "外部当前项目",
            "external_company": "外部公司",
            "role": "审校",
            "start_date": "2026-06-01",
            "remaining_volume": 6000,
            "deadline": "2026-08-15",
        },
        token=ET,
    )
    _, past_external = req(
        "POST", f"/api/translators/{experience_tid}/project-experiences",
        {
            "cooperation_source": "external",
            "project_status": "past",
            "project_name": "外部过往项目",
            "role": "LQE",
            "start_date": "2025-01-01",
            "end_date": "2025-12-01",
        },
        token=ET,
    )
    _, experiences = req(
        "GET", f"/api/translators/{experience_tid}/project-experiences",
    )
    chk("一个译员可新增多条项目经历且返回稳定ID",
        len(experiences) == 3
        and all(isinstance(item.get("id"), int) for item in experiences),
        experiences)
    chk("项目经历按当前优先和开始日期倒序",
        [item["project_name"] for item in experiences]
        == ["我司当前项目", "外部当前项目", "外部过往项目"],
        experiences)
    chk("项目经历语言码复用现有校验并规范化",
        current_our.get("source_lang") == "ZH"
        and current_our.get("target_lang") == "EN",
        current_our)
    chk("同一译员支持两个当前项目和独立角色",
        current_our.get("role") == "翻译"
        and current_external.get("role") == "审校"
        and past_external.get("project_status") == "past",
        experiences)
    profile_after_experiences = next(
        item for item in req("GET", "/api/translators")[1]
        if item["id"] == experience_tid
    )
    chk("项目经历读写不改变旧当前项目和角色",
        profile_after_experiences.get("current_project") == "旧字段项目"
        and profile_after_experiences.get("role") == "翻译",
        profile_after_experiences)
    chk("译员列表返回全部当前项目摘要",
        [item["id"] for item in profile_after_experiences.get("current_projects", [])]
        == [current_our["id"], current_external["id"]],
        profile_after_experiences.get("current_projects"))
    _, updated_experience = req(
        "PUT",
        f"/api/translators/{experience_tid}/project-experiences/{current_our['id']}",
        {
            "cooperation_source": "our_company",
            "project_status": "current",
            "project_name": "我司当前项目（更新）",
            "role": "LQA",
            "source_lang": "ZH",
            "target_lang": "EN",
            "start_date": "2026-07-01",
            "remaining_volume": 9000,
            "deadline": "2026-08-01",
        },
        token=ET,
    )
    chk("项目经历可编辑并保持稳定ID",
        updated_experience.get("id") == current_our["id"]
        and updated_experience.get("project_name") == "我司当前项目（更新）"
        and updated_experience.get("role") == "LQA"
        and updated_experience.get("remaining_volume") == 9000,
        updated_experience)
    req(
        "DELETE",
        f"/api/translators/{experience_tid}/project-experiences/{past_external['id']}",
        token=ET,
    )
    _, experiences_after_delete = req(
        "GET", f"/api/translators/{experience_tid}/project-experiences",
    )
    chk("项目经历可删除且不影响其他行",
        {item["id"] for item in experiences_after_delete}
        == {current_our["id"], current_external["id"]},
        experiences_after_delete)
    export_bytes = req("GET", "/api/export/translators")[1]
    export_workbook = load_workbook(io.BytesIO(export_bytes), data_only=True)
    project_export_rows = list(
        export_workbook["项目经历"].iter_rows(values_only=True)
    )
    chk("译员导出包含项目经历工作表",
        project_export_rows[0][0:4]
        == ("id", "translator_id", "translator_name", "translator_email")
        and any(row[6] == "我司当前项目（更新）" for row in project_export_rows[1:]),
        project_export_rows[:3])
    alias_export_rows = list(
        export_workbook["名称映射"].iter_rows(values_only=True)
    )
    chk(
        "译员导出包含稳定ID名称映射工作表",
        alias_export_rows[0]
        == (
            "id", "translator_id", "translator_name",
            "translator_email", "alias",
        )
        and any(
            row[1] == imported_profile["id"] and row[4] == "导入新"
            for row in alias_export_rows[1:]
        ),
        alias_export_rows[:5],
    )

    imported_project_email = f"project-import{run_tag}@x.com"
    workbook_raw, workbook_ct = mp(mk_xlsx_sheets({
        "译员": [
            ["姓名", "邮箱", "母语", "入库日期", "性别", "主体类型"],
            ["项目导入译员" + run_tag, imported_project_email, "英语", "2026-07-20", "男", "供应商"],
        ],
        "项目经历": [
            ["译员邮箱", "合作来源", "项目状态", "项目名称", "合作公司", "角色", "源语言", "目标语言"],
            [imported_project_email, "与别家合作", "过往项目", "导入项目", "外部工作室", "翻译", "EN", "ZH-HANS"],
        ],
    }))
    _, workbook_import = req(
        "POST", "/api/import/translators",
        raw=workbook_raw, token=ET, ct=workbook_ct,
    )
    imported_project_profile = next(
        item for item in req("GET", "/api/translators")[1]
        if item.get("email") == imported_project_email
    )
    imported_project_rows = req(
        "GET",
        f"/api/translators/{imported_project_profile['id']}/project-experiences",
    )[1]
    chk("同一工作簿可导入译员及项目经历",
        workbook_import.get("imported") == 1
        and workbook_import.get("imported_projects") == 1
        and imported_project_profile.get("gender") == "male"
        and imported_project_profile.get("entity_type") == "vendor"
        and imported_project_rows[0].get("project_name") == "导入项目"
        and imported_project_rows[0].get("cooperation_source") == "external",
        {"result": workbook_import, "projects": imported_project_rows})
    chk("不存在译员的项目经历GET返回404",
        code(lambda: req("GET", "/api/translators/99999/project-experiences")) == 404)
    chk("不存在译员的项目经历POST返回404",
        code(lambda: req(
            "POST", "/api/translators/99999/project-experiences",
            {
                "cooperation_source": "our_company",
                "project_status": "current",
                "project_name": "不存在",
                "start_date": "2026-08-01",
                "remaining_volume": 1000,
                "deadline": "2026-08-31",
            },
            token=ET,
        )) == 404)
    chk("只读角色不能新增项目经历",
        code(lambda: req(
            "POST", f"/api/translators/{experience_tid}/project-experiences",
            {
                "cooperation_source": "our_company",
                "project_status": "current",
                "project_name": "只读写入",
                "start_date": "2026-08-01",
                "remaining_volume": 1000,
                "deadline": "2026-08-31",
            },
            token=BT,
        )) == 403)
    chk("agent不能新增项目经历",
        code(lambda: req(
            "POST", f"/api/translators/{experience_tid}/project-experiences",
            {
                "cooperation_source": "our_company",
                "project_status": "current",
                "project_name": "Agent写入",
                "start_date": "2026-08-01",
                "remaining_volume": 1000,
                "deadline": "2026-08-31",
            },
            token=AT,
        )) == 403)
    chk("只读角色不能编辑项目经历",
        code(lambda: req(
            "PUT",
            f"/api/translators/{experience_tid}/project-experiences/{current_our['id']}",
            {
                "cooperation_source": "our_company",
                "project_status": "current",
                "project_name": "只读编辑",
                "start_date": "2026-08-01",
                "remaining_volume": 1000,
                "deadline": "2026-08-31",
            },
            token=BT,
        )) == 403)
    chk("agent不能删除项目经历",
        code(lambda: req(
            "DELETE",
            f"/api/translators/{experience_tid}/project-experiences/{current_our['id']}",
            token=AT,
        )) == 403)
    chk("跨译员编辑项目经历返回404",
        code(lambda: req(
            "PUT",
            f"/api/translators/1/project-experiences/{current_our['id']}",
            {
                "cooperation_source": "our_company",
                "project_status": "current",
                "project_name": "跨译员",
                "start_date": "2026-08-01",
                "remaining_volume": 1000,
                "deadline": "2026-08-31",
            },
            token=ET,
        )) == 404)
    chk("非法项目状态返回422",
        code(lambda: req(
            "POST", f"/api/translators/{experience_tid}/project-experiences",
            {
                "cooperation_source": "our_company",
                "project_status": "active",
                "project_name": "非法状态",
            },
            token=ET,
        )) == 422)
    chk("非法合作来源返回422",
        code(lambda: req(
            "POST", f"/api/translators/{experience_tid}/project-experiences",
            {
                "cooperation_source": "partner",
                "project_status": "current",
                "project_name": "非法来源",
            },
            token=ET,
        )) == 422)
    chk("负数剩余量返回422",
        code(lambda: req(
            "POST", f"/api/translators/{experience_tid}/project-experiences",
            {
                "cooperation_source": "our_company",
                "project_status": "current",
                "project_name": "负数剩余量",
                "remaining_volume": -1,
            },
            token=ET,
        )) == 422)
    chk("语言只填一侧返回422",
        code(lambda: req(
            "POST", f"/api/translators/{experience_tid}/project-experiences",
            {
                "cooperation_source": "our_company",
                "project_status": "current",
                "project_name": "缺目标语言",
                "source_lang": "ZH",
            },
            token=ET,
        )) == 422)
    chk("非法项目经历语言码返回400",
        code(lambda: req(
            "POST", f"/api/translators/{experience_tid}/project-experiences",
            {
                "cooperation_source": "our_company",
                "project_status": "current",
                "project_name": "非法语言码",
                "source_lang": "XX",
                "target_lang": "EN",
                "start_date": "2026-08-01",
                "remaining_volume": 1000,
                "deadline": "2026-08-31",
            },
            token=ET,
        )) == 400)
    _, audit_rows = req("GET", "/api/audit", token=ET)
    chk("新增项目经历写入审计日志",
        any(
            item.get("entity") == "项目经历"
            and item.get("entity_id") == current_our.get("id")
            and "我司当前项目" in (item.get("detail") or "")
            for item in audit_rows
        ),
        audit_rows[:5])
    chk("编辑和删除项目经历写入审计日志",
        any(
            item.get("action") == "编辑"
            and item.get("entity") == "项目经历"
            and item.get("entity_id") == current_our.get("id")
            for item in audit_rows
        )
        and any(
            item.get("action") == "删除"
            and item.get("entity") == "项目经历"
            and item.get("entity_id") == past_external.get("id")
            for item in audit_rows
        ),
        audit_rows[:10])

    print("=== D 安全（加密脱敏）===")
    _, pm = req("GET", "/api/translators/1/payment")
    chk("银行账号脱敏", pm["bank_account"].startswith("*") and pm["bank_account"].endswith("5678"))
    _, pf = req("GET", "/api/translators/1/payment/reveal", token=ET)
    chk("editor看明文全号", pf["bank_account"] == "6225888812345678")
    chk("viewer看明文403", code(lambda: req("GET", "/api/translators/1/payment/reveal", token=BT)) == 403)

    print("=== E agent 护栏 ===")
    pending_before = len(req("GET", "/api/pending", token=ET)[1])
    safe_rate = {"change_date": "2026-06-24", "source_lang": "ZH", "target_lang": "EN",
                 "task_type": "翻译", "new_rate": 771}
    safety_headers = {"Idempotency-Key": "acceptance:" + run_tag + ":rate",
                      "X-Dry-Run": "true"}
    _, rate_preview = req("POST", "/api/translators/1/rate-changes", safe_rate,
                          token=AT, headers=safety_headers)
    chk("agent费率dry-run不落待审",
        rate_preview.get("dry_run") is True and len(req("GET", "/api/pending", token=ET)[1]) == pending_before,
        rate_preview)

    safety_headers.pop("X-Dry-Run")
    _, safe_first = req("POST", "/api/translators/1/rate-changes", safe_rate,
                        token=AT, headers=safety_headers)
    _, safe_retry = req("POST", "/api/translators/1/rate-changes", safe_rate,
                        token=AT, headers=safety_headers)
    chk("agent同幂等键重试返回原待审",
        safe_retry.get("pending_id") == safe_first.get("pending_id") and safe_retry.get("deduplicated") is True,
        safe_retry)
    _, equivalent_same_key = req(
        "POST", "/api/translators/1/rate-changes",
        dict(safe_rate, source_lang="zh", target_lang="en"), token=AT, headers=safety_headers,
    )
    chk("agent同键业务等价内容返回原待审",
        equivalent_same_key.get("pending_id") == safe_first.get("pending_id")
        and equivalent_same_key.get("deduplicated") is True,
        equivalent_same_key)
    changed_rate = dict(safe_rate, new_rate=772)
    chk("agent同幂等键不同内容409",
        code(lambda: req("POST", "/api/translators/1/rate-changes", changed_rate,
                         token=AT, headers=safety_headers)) == 409)
    equivalent_rate = dict(safe_rate, source_lang=" zh ", target_lang="en")
    _, content_retry = req("POST", "/api/translators/1/rate-changes", equivalent_rate, token=AT,
                           headers={"Idempotency-Key": "acceptance:" + run_tag + ":rate-copy"})
    chk("agent规范化后相同活动内容防重",
        content_retry.get("pending_id") == safe_first.get("pending_id") and content_retry.get("deduplicated") is True,
        content_retry)
    req("POST", f"/api/pending/{safe_first['pending_id']}/reject", token=ET)
    _, safe_resubmit = req("POST", "/api/translators/1/rate-changes", safe_rate, token=AT,
                           headers={"Idempotency-Key": "acceptance:" + run_tag + ":rate-resubmit"})
    chk("已处理内容可用新幂等键重提", safe_resubmit.get("pending_id") != safe_first.get("pending_id"), safe_resubmit)
    req("POST", f"/api/pending/{safe_resubmit['pending_id']}/reject", token=ET)

    _, po_preview = req("POST", "/api/po",
                        {"translator_id": 1, "settlement_month": "2026-10", "role": "翻译", "word_count": 1000},
                        token=AT, headers={"X-Dry-Run": "true"})
    chk("agent PO dry-run不落库",
        po_preview.get("dry_run") is True
        and len(req("GET", "/api/po?month=2026-10")[1]) == 0
        and len(req("GET", "/api/pending", token=ET)[1]) == 0,
        po_preview)
    _, status_preview = req("PUT", "/api/po/1/status", {"status": "已支付"}, token=AT,
                            headers={"X-Dry-Run": "true"})
    po_status_after_preview = next(p["status"] for p in req("GET", "/api/po")[1] if p["id"] == 1)
    chk("agent PO状态dry-run不落库",
        status_preview.get("dry_run") is True and po_status_after_preview == "已开票待付"
        and len(req("GET", "/api/pending", token=ET)[1]) == 0,
        status_preview)
    chk("agent费率dry-run校验语言对",
        code(lambda: req("POST", "/api/translators/1/rate-changes",
                         dict(safe_rate, source_lang="xx", target_lang="KO"), token=AT,
                         headers={"X-Dry-Run": "true"})) == 400)
    chk("agent PO dry-run校验译员",
        code(lambda: req("POST", "/api/po",
                         {"translator_id": 99999, "settlement_month": "2026-10"}, token=AT,
                         headers={"X-Dry-Run": "true"})) == 404)
    chk("agent PO状态dry-run校验目标PO",
        code(lambda: req("PUT", "/api/po/99999/status", {"status": "已支付"}, token=AT,
                         headers={"X-Dry-Run": "true"})) == 404)
    durable_po = {"translator_id": 1, "settlement_month": "2026-12", "role": "翻译",
                  "word_count": 1000, "rate": 180, "po_number": "PO-IDEMPOTENT-" + run_tag}
    durable_headers = {"Idempotency-Key": "acceptance:" + run_tag + ":durable-po"}
    _, durable_pending = req("POST", "/api/po", durable_po, token=AT, headers=durable_headers)
    req("POST", f"/api/pending/{durable_pending['pending_id']}/approve", token=ET)
    try:
        _, durable_retry = req("POST", "/api/po", durable_po, token=AT, headers=durable_headers)
    except urllib.error.HTTPError as exc:
        durable_retry = {"http_error": exc.code}
    chk("已批准PO同键重试返回原结果",
        durable_retry.get("pending_id") == durable_pending.get("pending_id")
        and durable_retry.get("status") == "approved" and durable_retry.get("deduplicated") is True,
        durable_retry)
    chk("已批准PO同键不同请求409",
        code(lambda: req("POST", "/api/po", dict(durable_po, settlement_month="2027-01"),
                         token=AT, headers=durable_headers)) == 409)

    before_agent_rate = [t for t in req("GET", "/api/translators")[1] if t["id"] == 1][0]["translation_rate"]
    _, pr = req("POST", "/api/translators/1/rate-changes", {"change_date": "2026-06-25", "task_type": "翻译", "new_rate": 777}, token=AT)
    chk("agent费率→待审", pr.get("pending") is True)
    chk("待审期间主表未变", [t for t in req("GET", "/api/translators")[1] if t["id"] == 1][0]["translation_rate"] == before_agent_rate)
    req("POST", f"/api/pending/{pr['pending_id']}/reject", token=ET)
    chk("驳回后待审清空", len(req("GET", "/api/pending", token=ET)[1]) == 0)
    req(
        "PUT",
        "/api/translators/1/capacity/override?month=2026-08",
        {"status": "健康", "reason": "Agent 已人工确认"},
        token=AT,
    )
    agent_capacity = req(
        "GET", "/api/translators/1/capacity?month=2026-08",
    )[1]
    chk(
        "agent可按月修正档期且不录入百分比",
        agent_capacity.get("capacity_override", {}).get("status") == "健康"
        and "occupancy_pct" not in json.dumps(agent_capacity, ensure_ascii=False),
        agent_capacity,
    )

    print("=== F 优化校验（枚举/唯一性）===")
    chk("非法译员状态422", code(lambda: req("POST", "/api/translators", {"name": "状态X", "status": "乱", "native_language": "中文", "onboarding_date": "2025-01-01"}, token=ET)) == 422)
    chk("非法入库日期422", code(lambda: req("POST", "/api/translators", {"name": "日期X", "native_language": "中文", "onboarding_date": "2025-01-32"}, token=ET)) == 422)
    chk("非法邮箱422", code(lambda: req("POST", "/api/translators", {"name": "邮箱X", "email": "bad-email", "native_language": "中文", "onboarding_date": "2025-01-01"}, token=ET)) == 422)
    chk("非法PO结算月422", code(lambda: req("POST", "/api/po", {"translator_id": 1, "settlement_month": "2026-13"}, token=ET)) == 422)
    chk("非法产能月份400", code(lambda: req("GET", "/api/translators/1/capacity?month=2026-13")) == 400)
    chk("非法月度修正状态422", code(lambda: req(
        "PUT", "/api/translators/1/capacity/override?month=2026-08",
        {"status": "满负荷", "reason": "测试"}, token=ET,
    )) == 422)
    chk("月度修正缺原因422", code(lambda: req(
        "PUT", "/api/translators/1/capacity/override?month=2026-08",
        {"status": "健康", "reason": ""}, token=ET,
    )) == 422)
    chk("非法客诉严重度422", code(lambda: req("POST", "/api/translators/1/complaints", {"severity": "乱"}, token=ET)) == 422)
    chk("重复合同编号400", code(lambda: req("POST", "/api/translators/1/contracts", {"contract_number": "C-2025-011"}, token=ET)) == 400)
    chk("重复PO号400", code(lambda: req("POST", "/api/po", {"translator_id": 1, "settlement_month": "2026-07", "po_number": "PO-202606-001"}, token=ET)) == 400)

    print("=== G 主表联动（PRD §6.3.3/§12.3.2）===")
    def tr(i): return [t for t in req("GET", "/api/translators")[1] if t["id"] == i][0]
    b3 = tr(3)
    req("POST", "/api/translators/3/complaints", {"date": "2026-06-20", "complaint_type": "质量不达标", "severity": "一般", "deduction_amount": 150}, token=ET)
    a3 = tr(3)
    chk("客诉联动 次数+1", a3["complaint_count"] == (b3["complaint_count"] or 0) + 1)
    chk("客诉联动 扣款累加150", round(a3["deduction_total"] - (b3["deduction_total"] or 0), 2) == 150)
    req("POST", "/api/translators/3/quality", {"evaluation_period": "2026-06", "qa_type": "LQE", "score": 88, "minor_errors": 2}, token=ET)
    req("POST", "/api/translators/3/quality", {"evaluation_period": "2026-07", "qa_type": "LQE", "score": 92, "minor_errors": 1}, token=ET)
    q3 = tr(3)
    chk("质量联动 最近分=92", q3["recent_qa_score"] == 92)
    chk("质量联动 累计均分=90", q3["cumulative_qa_score"] == 90)
    chk("累计均分90自动评级A", q3["internal_rating"] == "A", q3)
    chk("质量联动 低错累计+3", q3["low_error_count"] == (b3["low_error_count"] or 0) + 3)
    chk("谈判联动 成功后谈后费率", (lambda t: (req("POST", "/api/translators/3/rate-changes", {"change_date": "2026-06-25", "task_type": "翻译", "original_rate": 150, "new_rate": 140, "result": "成功"}, token=ET), tr(3))[1]["post_negotiation_rate"] == 140)(None))

    print("=== H 子表删除 ===")
    _, del_tr = req("POST", "/api/translators", {
        "name": "删除校验" + run_tag,
        "email": f"delete{run_tag}@x.com",
        "native_language": "中文",
        "onboarding_date": "2026-07-01",
        "gender": "male",
    }, token=ET)
    did = del_tr["id"]
    req("POST", f"/api/translators/{did}/language-pairs",
        {"source_lang": "ZH", "target_lang": "KO", "translation_rate": 200}, token=ET)
    lp_id = req("GET", f"/api/translators/{did}/language-pairs")[1][0]["id"]
    req("DELETE", f"/api/translators/{did}/language-pairs/{lp_id}", token=ET)
    chk("删除语言对后列表为空", req("GET", f"/api/translators/{did}/language-pairs")[1] == [])
    req("POST", f"/api/translators/{did}/rate-changes",
        {"change_date": "2026-07-01", "task_type": "翻译", "original_rate": 200, "new_rate": 180, "result": "成功"},
        token=ET)
    rc_id = req("GET", f"/api/translators/{did}/rate-changes")[1][0]["id"]
    req("DELETE", f"/api/translators/{did}/rate-changes/{rc_id}", token=ET)
    chk("删除报价变更后列表为空", req("GET", f"/api/translators/{did}/rate-changes")[1] == [])
    req("POST", f"/api/translators/{did}/quality", {"evaluation_period": "2026-06", "qa_type": "LQE", "score": 80, "minor_errors": 2}, token=ET)
    req("POST", f"/api/translators/{did}/quality", {"evaluation_period": "2026-07", "qa_type": "LQE", "score": 100, "minor_errors": 1}, token=ET)
    q_id = req("GET", f"/api/translators/{did}/quality")[1][0]["id"]
    req("DELETE", f"/api/translators/{did}/quality/{q_id}", token=ET)
    dq = tr(did)
    chk("删除质量后重算最近分", dq["recent_qa_score"] == 80 and dq["cumulative_qa_score"] == 80 and dq["low_error_count"] == 2, dq)
    chk("删除质量后自动重算评级B", dq["internal_rating"] == "B", dq)
    req("POST", f"/api/translators/{did}/contracts", {"contract_number": "DEL-" + run_tag, "status": "有效"}, token=ET)
    ct_id = req("GET", f"/api/translators/{did}/contracts")[1][0]["id"]
    req("DELETE", f"/api/translators/{did}/contracts/{ct_id}", token=ET)
    chk("删除合同后列表为空", req("GET", f"/api/translators/{did}/contracts")[1] == [])
    req("POST", f"/api/translators/{did}/complaints", {"severity": "一般", "deduction_amount": 120}, token=ET)
    cp_id = req("GET", f"/api/translators/{did}/complaints")[1][0]["id"]
    req("DELETE", f"/api/translators/{did}/complaints/{cp_id}", token=ET)
    dc = tr(did)
    chk("删除客诉后重算次数扣款", dc["complaint_count"] == 0 and dc["deduction_total"] == 0, dc)
    req(
        "PUT",
        f"/api/translators/{did}/capacity/override?month=2026-08",
        {"status": "空闲", "reason": "删除回归"},
        token=ET,
    )
    req(
        "DELETE",
        f"/api/translators/{did}/capacity/override?month=2026-08",
        token=ET,
    )
    chk(
        "清除按月人工修正后恢复自动状态",
        req("GET", f"/api/translators/{did}/capacity?month=2026-08")[1]["capacity_override"] is None,
    )
    req("PUT", f"/api/translators/{did}/payment",
        {"currency": "CNY", "bank_name": "测试银行", "bank_account": "6222000011112222", "payee_name": "删除校验"},
        token=ET)
    req("DELETE", f"/api/translators/{did}/payment", token=ET)
    chk("删除支付信息后为空", req("GET", f"/api/translators/{did}/payment")[1] is None)
    chk("boss删除子表403", code(lambda: req("DELETE", f"/api/translators/{did}/quality/1", token=BT)) == 403)

    print("=== I 反馈文档 13 项新增闭环 ===")
    filter_body = {
        "name": "精确筛选" + run_tag,
        "email": f"filter{run_tag}@x.com",
        "native_language": "English",
        "onboarding_date": "2026-07-29",
        "wechat": "filter_wechat",
        "domains": "RPG, SLG",
        "gender": "female",
        "entity_type": "individual",
        "daily_output": 3000,
        "settlement_mode": "cumulative",
    }
    _, filter_translator = req(
        "POST", "/api/translators", filter_body, token=ET,
    )
    filter_tid = filter_translator["id"]
    chk(
        "新增译员尚无 LQE 时不自动评级",
        filter_translator.get("internal_rating") is None,
        filter_translator,
    )
    req(
        "POST", f"/api/translators/{filter_tid}/quality",
        {"evaluation_period": "2026-07", "qa_type": "LQE", "score": 87},
        token=ET,
    )
    auto_rated = tr(filter_tid)
    chk(
        "累计均分87自动评级A-",
        auto_rated.get("cumulative_qa_score") == 87
        and auto_rated.get("internal_rating") == "A-",
        auto_rated,
    )
    _, manual_update = req(
        "PUT",
        f"/api/translators/{filter_tid}",
        filter_body | {
            "manual_rating": "S",
            "manual_rating_reason": "非质量原因暂停合作",
        },
        token=ET,
    )
    chk(
        "人工例外评级可覆盖 LQE 自动评级并记录原因",
        manual_update.get("internal_rating") == "S"
        and manual_update.get("manual_rating") == "S"
        and manual_update.get("manual_rating_reason") == "非质量原因暂停合作",
        manual_update,
    )
    _, cleared_manual = req(
        "PUT",
        f"/api/translators/{filter_tid}",
        filter_body | {"manual_rating": None, "manual_rating_reason": None},
        token=ET,
    )
    chk(
        "清除人工例外评级后恢复 LQE 自动评级",
        cleared_manual.get("internal_rating") == "A-",
        cleared_manual,
    )
    req(
        "POST", f"/api/translators/{filter_tid}/language-pairs",
        {
            "source_lang": "ZH",
            "target_lang": "EN",
            "translation_rate": 200,
            "lqa_rate": 15,
            "currency": "CNY",
        },
        token=ET,
    )
    next_year = date.today().year + (1 if date.today().month == 12 else 0)
    next_month = 1 if date.today().month == 12 else date.today().month + 1
    capacity_month = f"{next_year:04d}-{next_month:02d}"
    capacity_start = date(next_year, next_month, 1).isoformat()
    deadline = date(
        next_year, next_month, monthrange(next_year, next_month)[1],
    ).isoformat()
    req(
        "POST", f"/api/translators/{filter_tid}/project-experiences",
        {
            "cooperation_source": "our_company",
            "project_status": "current",
            "project_name": "高负荷翻译",
            "role": "翻译",
            "source_lang": "ZH",
            "target_lang": "EN",
            "start_date": capacity_start,
            "remaining_volume": 300000,
            "deadline": deadline,
        },
        token=ET,
    )
    req(
        "POST", f"/api/translators/{filter_tid}/project-experiences",
        {
            "cooperation_source": "our_company",
            "project_status": "current",
            "project_name": "LQA 小时",
            "role": "LQA",
            "source_lang": "ZH",
            "target_lang": "EN",
            "start_date": capacity_start,
            "remaining_volume": 30,
            "deadline": deadline,
        },
        token=ET,
    )
    req(
        "PUT",
        f"/api/translators/{filter_tid}/capacity/override?month={capacity_month}",
        {"status": "空闲", "reason": "人工确认本月不再接新项目"},
        token=ET,
    )
    computed = next(
        item for item in req(
            "GET", f"/api/translators?capacity_month={capacity_month}",
        )[1]
        if item["id"] == filter_tid
    )
    chk(
        "档期按月度项目量和日产能计算并提示人工冲突",
        computed.get("computed_availability") == "警告"
        and computed.get("effective_availability") == "空闲"
        and computed.get("capacity_month") == capacity_month
        and computed.get("availability_conflict") is True
        and len(computed.get("availability_basis", [])) == 2,
        computed,
    )
    basis_by_role = {
        item["role"]: item for item in computed.get("availability_basis", [])
    }
    chk(
        "各任务沿用翻译日产能并按20个工作日计算月产能",
        basis_by_role["翻译"]["daily_capacity"] == 3000
        and basis_by_role["翻译"]["monthly_capacity"] == 60000
        and basis_by_role["LQA"]["daily_capacity"] == 3000
        and basis_by_role["LQA"]["monthly_capacity"] == 60000,
        basis_by_role,
    )
    monthly_capacity = req(
        "GET",
        f"/api/translators/{filter_tid}/capacity?month={capacity_month}",
    )[1]
    chk(
        "指定月份返回逐项目字数分摊且无手填占用百分比",
        monthly_capacity.get("allocated_words") == 300030
        and monthly_capacity.get("monthly_capacity") == 60000
        and monthly_capacity.get("capacity_data_complete") is True
        and all("allocated_words" in row for row in monthly_capacity.get("availability_basis", []))
        and "occupancy_pct" not in json.dumps(monthly_capacity, ensure_ascii=False),
        monthly_capacity,
    )

    query = urlencode({
        "source_lang": "ZH",
        "target_lang": "EN",
        "native_language": "English",
        "rate_type": "translation",
        "min_rate": 150,
        "max_rate": 300,
        "currency": "CNY",
        "domain": "RPG",
        "has_wechat": "true",
        "rating": "A-",
        "gender": "female",
        "entity_type": "individual",
        "paged": "true",
        "page": 1,
        "page_size": 20,
    })
    _, filtered = req("GET", "/api/translators?" + query)
    chk(
        "精确筛选全部条件按AND命中并返回分页总数",
        filtered.get("total") == 1
        and filtered["items"][0]["id"] == filter_tid,
        filtered,
    )
    _, overview = req("GET", "/api/overview")
    chk(
        "汇总面板返回母语、语言对和S/A/A-/B评级",
        overview["native_languages"].get("English") == 1
        and overview["language_pairs"].get("ZH→EN", 0) >= 1
        and overview["ratings"].get("A-", 0) >= 1
        and set(overview["ratings"]) == {"S", "A", "A-", "B"},
        overview,
    )
    lqe_raw, lqe_ct = mp(mk_xlsx([
        ["子表", "段数", "词数", "错误数", "严重错误", "SCORE", "STATUS"],
        [filter_translator["name"], 1, 100, 0, 0, 99, "PASS"],
    ]))
    _, lqe_import = req(
        "POST",
        "/api/import/lqe?period=2026-08",
        raw=lqe_raw,
        token=ET,
        ct=lqe_ct,
    )
    lqe_rated = tr(filter_tid)
    chk(
        "LQE 导入后累计均分与评级自动重算",
        lqe_import.get("imported") == 1
        and lqe_rated.get("cumulative_qa_score") == 93
        and lqe_rated.get("internal_rating") == "A",
        {"import": lqe_import, "translator": lqe_rated},
    )

    _, fixed_price = req(
        "POST", f"/api/translators/{filter_tid}/project-prices",
        {
            "project_name": "整包项目",
            "price_type": "fixed",
            "amount": 1200,
            "unit": "project",
            "currency": "CNY",
        },
        token=ET,
    )
    _, custom_price = req(
        "POST", f"/api/translators/{filter_tid}/project-prices",
        {
            "project_name": "音频校验",
            "price_type": "custom",
            "custom_task_name": "音频逐条检查",
            "amount": 80,
            "unit": "hour",
            "currency": "CNY",
        },
        token=ET,
    )
    _, translation_project_price = req(
        "POST", f"/api/translators/{filter_tid}/project-prices",
        {
            "project_name": "角色项目",
            "source_lang": "ZH",
            "target_lang": "EN",
            "price_type": "translation",
            "amount": 260,
            "unit": "per_1000",
            "currency": "CNY",
        },
        token=ET,
    )
    _, review_project_price = req(
        "POST", f"/api/translators/{filter_tid}/project-prices",
        {
            "project_name": "角色项目",
            "source_lang": "ZH",
            "target_lang": "EN",
            "price_type": "review",
            "amount": 160,
            "unit": "per_1000",
            "currency": "CNY",
        },
        token=ET,
    )
    chk(
        "翻译、审校、一口价和自定义其他价格可独立保存",
        fixed_price.get("task_type") == "一口价"
        and fixed_price.get("amount") == 1200
        and custom_price.get("custom_task_name") == "音频逐条检查"
        and translation_project_price.get("task_type") == "翻译"
        and review_project_price.get("task_type") == "审校",
        {
            "translation": translation_project_price,
            "review": review_project_price,
            "fixed": fixed_price,
            "custom": custom_price,
        },
    )
    _, updated_fixed_price = req(
        "PUT",
        f"/api/translators/{filter_tid}/project-prices/{fixed_price['id']}",
        {
            "project_name": "整包项目",
            "price_type": "fixed",
            "amount": 1300,
            "unit": "project",
            "currency": "CNY",
        },
        token=ET,
    )
    chk(
        "项目专用价格可编辑",
        updated_fixed_price.get("amount") == 1300,
        updated_fixed_price,
    )
    generic_filters = json.dumps([
        {
            "field": "name",
            "op": "eq",
            "value": filter_translator["name"],
        },
        {"field": "gender", "op": "eq", "value": "female"},
        {
            "field": "project_prices.amount",
            "op": "between",
            "value": "1200,1400",
        },
    ], ensure_ascii=False)
    _, generic_filtered = req(
        "GET",
        "/api/translators?" + urlencode({
            "filters": generic_filters,
            "paged": "true",
            "page": 1,
            "page_size": 20,
        }),
    )
    chk(
        "主表与业务子表所有字段可组合筛选",
        generic_filtered.get("total") == 1
        and generic_filtered["items"][0]["id"] == filter_tid,
        generic_filtered,
    )
    _, price_match = req(
        "GET",
        "/api/po/price-match?" + urlencode({
            "translator_id": filter_tid,
            "project": "整包项目",
            "role": "一口价",
            "currency": "CNY",
        }),
    )
    _, auto_fixed_po = req(
        "POST", "/api/po",
        {
            "translator_id": filter_tid,
            "settlement_month": "2026-10",
            "project": "整包项目",
            "role": "一口价",
            "currency": "CNY",
            "po_number": "PO-AUTO-FIXED-" + run_tag,
        },
        token=ET,
    )
    chk(
        "PO按最新项目价格自动匹配且保留来源",
        price_match.get("matched") is True
        and price_match.get("project_price_id") == fixed_price["id"]
        and price_match.get("amount") == 1300
        and auto_fixed_po.get("pricing_mode") == "fixed"
        and auto_fixed_po.get("amount") == 1300,
        {"match": price_match, "po": auto_fixed_po},
    )
    _, translation_price_match = req(
        "GET",
        "/api/po/price-match?" + urlencode({
            "translator_id": filter_tid,
            "project": "角色项目",
            "role": "翻译",
            "source_lang": "ZH",
            "target_lang": "EN",
            "currency": "CNY",
        }),
    )
    _, review_price_match = req(
        "GET",
        "/api/po/price-match?" + urlencode({
            "translator_id": filter_tid,
            "project": "角色项目",
            "role": "审校",
            "source_lang": "ZH",
            "target_lang": "EN",
            "currency": "CNY",
        }),
    )
    chk(
        "PO优先匹配翻译和审校项目价格",
        translation_price_match.get("source") == "project_price"
        and translation_price_match.get("project_price_id") == translation_project_price["id"]
        and translation_price_match.get("rate") == 260
        and review_price_match.get("source") == "project_price"
        and review_price_match.get("project_price_id") == review_project_price["id"]
        and review_price_match.get("rate") == 160,
        {"translation": translation_price_match, "review": review_price_match},
    )
    chk(
        "自定义其他价格缺任务名称422",
        code(lambda: req(
            "POST", f"/api/translators/{filter_tid}/project-prices",
            {
                "project_name": "坏价格",
                "price_type": "custom",
                "amount": 1,
                "unit": "other",
                "currency": "CNY",
            },
            token=ET,
        )) == 422,
    )

    _, fixed_po = req(
        "POST", "/api/po",
        {
            "translator_id": filter_tid,
            "settlement_month": "2026-11",
            "project": "整包项目",
            "role": "一口价",
            "pricing_mode": "fixed",
            "amount": 1200,
            "currency": "CNY",
            "po_number": "PO-FIXED-" + run_tag,
        },
        token=ET,
    )
    _, hourly_po = req(
        "POST", "/api/po",
        {
            "translator_id": filter_tid,
            "settlement_month": "2026-11",
            "project": "LQA 小时",
            "role": "LQA",
            "pricing_mode": "per_hour",
            "word_count": 3,
            "rate": 15,
            "currency": "USD",
            "po_number": "PO-HOUR-" + run_tag,
        },
        token=ET,
    )
    _, manual_po = req(
        "POST", "/api/po",
        {
            "translator_id": filter_tid,
            "settlement_month": "2026-11",
            "project": "手工调整",
            "role": "其他",
            "pricing_mode": "manual",
            "amount": 99,
            "currency": "CNY",
            "po_number": "PO-MANUAL-" + run_tag,
        },
        token=ET,
    )
    chk(
        "PO固定价、按小时和手工金额分支正确",
        fixed_po["amount"] == 1200
        and hourly_po["amount"] == 45
        and manual_po["amount"] == 99,
        {"fixed": fixed_po, "hourly": hourly_po, "manual": manual_po},
    )
    _, corrected_po = req(
        "PUT", f"/api/po/{manual_po['id']}",
        {
            "translator_id": filter_tid,
            "settlement_month": "2026-11",
            "project": "手工调整",
            "role": "其他",
            "pricing_mode": "manual",
            "amount": 109,
            "currency": "CNY",
            "status": "有争议",
            "po_number": "PO-MANUAL-" + run_tag,
            "remarks": "人工复核修正",
        },
        token=ET,
    )
    _, unpaid_details = req("GET", f"/api/po/unpaid/{filter_tid}")
    chk(
        "PO支持人工修正并按译员展示未结算明细",
        corrected_po.get("amount") == 109
        and corrected_po.get("status") == "有争议"
        and all(
            row["status"] in {"未开票", "已开票待付"}
            for row in unpaid_details["rows"]
        )
        and unpaid_details["totals"].get("CNY", 0) >= 2500,
        {"corrected": corrected_po, "unpaid": unpaid_details},
    )
    po_filter_query = urlencode({
        "month": "2026-11",
        "translator": filter_translator["name"],
        "project": "整包",
        "status": "未开票",
        "role": "一口价",
        "currency": "cny",
        "pricing_mode": "fixed",
        "settlement_mode": "cumulative",
        "min_amount": 1000,
        "max_amount": 1500,
        "po_number": "FIXED-" + run_tag,
    })
    _, filtered_po = req("GET", "/api/po?" + po_filter_query)
    chk(
        "PO结算支持金额区间等十一项AND组合筛选",
        len(filtered_po) == 1 and filtered_po[0]["id"] == fixed_po["id"],
        filtered_po,
    )
    _, filtered_po_summary = req("GET", "/api/po/summary?" + po_filter_query)
    chk(
        "PO筛选条件同步作用于当月及跨月汇总",
        filtered_po_summary["by_currency"] == {
            "CNY": {"unpaid": 1200.0, "paid": 0.0},
        }
        and filtered_po_summary["cumulative_unpaid_by_currency"] == {
            "CNY": 2500.0,
        },
        filtered_po_summary,
    )
    mismatch_filter_query = urlencode({
        "month": "2026-11",
        "translator": filter_translator["name"],
        "project": "整包",
        "currency": "USD",
    })
    _, mismatched_po = req("GET", "/api/po?" + mismatch_filter_query)
    chk("PO组合筛选任一条件不符即不命中", mismatched_po == [], mismatched_po)

    projectlist_alias = "Projectlist别名" + run_tag
    _, created_alias = req(
        "POST",
        f"/api/translators/{filter_tid}/aliases",
        {"alias": projectlist_alias},
        token=ET,
    )
    _, alias_search = req(
        "GET",
        "/api/translators?" + urlencode({"q": projectlist_alias}),
    )
    chk(
        "名称映射可维护且可用于译员搜索",
        created_alias["alias"] == projectlist_alias
        and len(alias_search) == 1
        and alias_search[0]["id"] == filter_tid,
        {"alias": created_alias, "search": alias_search},
    )
    projectlist_alias = projectlist_alias + "更新"
    _, updated_alias = req(
        "PUT",
        f"/api/translators/{filter_tid}/aliases/{created_alias['id']}",
        {"alias": projectlist_alias},
        token=ET,
    )
    _, updated_alias_search = req(
        "GET",
        "/api/translators?" + urlencode({"q": projectlist_alias}),
    )
    chk(
        "名称映射可编辑且更新后仍指向同一稳定ID",
        updated_alias["id"] == created_alias["id"]
        and updated_alias["alias"] == projectlist_alias
        and len(updated_alias_search) == 1
        and updated_alias_search[0]["id"] == filter_tid,
        {"alias": updated_alias, "search": updated_alias_search},
    )
    _, temporary_alias = req(
        "POST",
        f"/api/translators/{filter_tid}/aliases",
        {"alias": "临时别名" + run_tag},
        token=ET,
    )
    req(
        "DELETE",
        f"/api/translators/{filter_tid}/aliases/{temporary_alias['id']}",
        token=ET,
    )
    remaining_aliases = req(
        "GET",
        f"/api/translators/{filter_tid}/aliases",
    )[1]
    chk(
        "名称映射可删除且不影响其他映射",
        all(item["id"] != temporary_alias["id"] for item in remaining_aliases)
        and any(item["id"] == created_alias["id"] for item in remaining_aliases),
        remaining_aliases,
    )

    projectlist_prefix = "Projectlist导入" + run_tag
    projectlist_projects = {
        "translation": projectlist_prefix + "-翻译",
        "review": projectlist_prefix + "-审校",
        "mtpe": projectlist_prefix + "-MTPE",
        "lqa": projectlist_prefix + "-LQA",
        "lqe": projectlist_prefix + "-LQE",
        "fixed": projectlist_prefix + "-一口价",
        "checked": projectlist_prefix + "-已勾选",
        "unsupported": projectlist_prefix + "-未知工作类型",
        "missing_wwc": projectlist_prefix + "-缺WWC",
    }
    projectlist_raw, projectlist_ct = mp(mk_xlsx_sheets({
        "【项目组】Projectlist (V1.0)": [
            ["说明"],
            [None],
            [
                "项目名称（稿件/LQA批次）", "DDL", "目标语言", "指定译员",
                "工作类型", "翻译费率", "币种", "REPNEW 实际",
                "译员WWC字数", "REPNEW & 小时数", "稿费金额（CNY）",
                "译员PO时间（X月）", "原语言", "结算PO", "已打款",
            ],
            [
                projectlist_projects["translation"], "2026-07-31", "英语",
                projectlist_alias, "翻译", 0.2, "CNY", 9900,
                1200, 77, 240, "7 月", "简体中文", "□", "否",
            ],
            [
                projectlist_projects["review"], "2026-07-31", "英语",
                filter_translator["name"], "审校", 0.1, "CNY", 8800,
                2000, 66, 200, "7 月", "简体中文", "□", "否",
            ],
            [
                projectlist_projects["mtpe"], "2026-07-31", "英语",
                filter_translator["name"], "MTPE", 0.08, "CNY", 7700,
                1500, 55, 120, "7 月", "简体中文", "□", "否",
            ],
            [
                projectlist_projects["lqa"], "2026-07-31", "藏语",
                filter_translator["name"], "LQA", 15, "USD", 6600,
                3, 66, None, "7 月", "藏语", "□", "否",
            ],
            [
                projectlist_projects["lqe"], "2026-07-31", "英语",
                filter_translator["name"], "LQE", 20, "USD", 5500,
                2, 55, None, "7 月", "简体中文", "□", "否",
            ],
            [
                projectlist_projects["fixed"], "2026-07-31", "英语",
                filter_translator["name"], "一口价", 999, "CNY", 4400,
                2200, 11, 680, "7 月", "简体中文", "□", "否",
            ],
            [
                projectlist_projects["checked"], "2026-07-31", "英语",
                filter_translator["name"], "翻译", 0.3, "CNY", 3300,
                1000, 22, None, "7 月", "简体中文", "✅", "否",
            ],
            [
                projectlist_projects["unsupported"], "2026-07-31", "英语",
                filter_translator["name"], "配音", 50, "CNY", 2200,
                1000, 4, None, "7 月", "简体中文", "□", "否",
            ],
            [
                projectlist_projects["missing_wwc"], "2026-07-31", "英语",
                filter_translator["name"], "审校", 0.1, "CNY", 2500,
                None, 6, None, "7 月", "简体中文", "□", "否",
            ],
        ],
    }))
    projectlist_secure_headers = [
        "项目名称（稿件/LQA批次）", "DDL", "目标语言", "指定译员",
        "工作类型", "翻译费率", "币种", "译员WWC字数",
        "稿费金额（CNY）", "译员PO时间（X月）", "原语言",
        "结算PO", "已打款",
    ]
    missing_projectlist_header_codes = {}
    for missing_header in ("结算PO", "已打款"):
        missing_header_raw, missing_header_ct = mp(mk_xlsx_sheets({
            "Projectlist 缺表头": [[
                header
                for header in projectlist_secure_headers
                if header != missing_header
            ]],
        }))
        missing_projectlist_header_codes[missing_header] = code(lambda: req(
            "POST", "/api/import/po",
            raw=missing_header_raw, token=ET, ct=missing_header_ct,
        ))
    chk(
        "Projectlist正式导入缺结算PO或已打款表头均拒绝",
        missing_projectlist_header_codes == {"结算PO": 400, "已打款": 400},
        missing_projectlist_header_codes,
    )
    _, projectlist_all = req(
        "POST", "/api/import/po?preview=true&projectlist_po_state=all",
        raw=projectlist_raw, token=ET, ct=projectlist_ct,
    )
    chk(
        "Projectlist预览统一取WWC并按工作类型计价",
        projectlist_all.get("source_format") == "projectlist"
        and projectlist_all.get("write_enabled") is True
        and projectlist_all.get("quantity_rule_confirmed") is True
        and projectlist_all.get("ready") == 6
        and projectlist_all.get("checked_rows") == 1
        and projectlist_all.get("unchecked_rows") == 8
        and projectlist_all.get("paid_rows") == 0
        and projectlist_all.get("settled_rows") == 1
        and projectlist_all.get("skipped_settled") == 1
        and len(projectlist_all.get("invalid_rows", [])) == 2
        and any(
            row.get("translator_id") == filter_tid
            and row.get("translator_name") == projectlist_alias
            and row.get("project") == projectlist_projects["translation"]
            and row.get("word_count") == 1200
            and row.get("rate") == 200
            and abs(float(row.get("expected_amount", 0)) - 240) <= 0.02
            and row.get("action") == "import"
            for row in projectlist_all.get("preview_rows", [])
        )
        and any(
            row.get("project") == projectlist_projects["lqa"]
            and row.get("pricing_mode") == "per_hour"
            and row.get("quantity_source") == "译员WWC字数"
            and row.get("word_count") == 3
            and abs(float(row.get("expected_amount", 0)) - 45) <= 0.02
            for row in projectlist_all.get("preview_rows", [])
        )
        and any(
            row.get("project") == projectlist_projects["fixed"]
            and row.get("pricing_mode") == "fixed"
            and abs(float(row.get("expected_amount", 0)) - 680) <= 0.02
            for row in projectlist_all.get("preview_rows", [])
        ),
        projectlist_all,
    )
    _, projectlist_checked = req(
        "POST", "/api/import/po?preview=true&projectlist_po_state=checked",
        raw=projectlist_raw, token=ET, ct=projectlist_ct,
    )
    _, projectlist_unchecked = req(
        "POST", "/api/import/po?preview=true&projectlist_po_state=unchecked",
        raw=projectlist_raw, token=ET, ct=projectlist_ct,
    )
    chk(
        "Projectlist支持已勾选和未勾选筛选",
        projectlist_checked.get("selected_rows") == 1
        and projectlist_checked.get("skipped_settled") == 1
        and projectlist_checked["preview_rows"][0]["action"] == "skip_settled"
        and projectlist_unchecked.get("selected_rows") == 8
        and projectlist_unchecked.get("ready") == 6
        and len(projectlist_unchecked.get("invalid_rows", [])) == 2
        and all(
            row["action"] == "import"
            for row in projectlist_unchecked.get("preview_rows", [])
        ),
        {
            "checked": projectlist_checked,
            "unchecked": projectlist_unchecked,
        },
    )
    _, projectlist_import = req(
        "POST", "/api/import/po?projectlist_po_state=all",
        raw=projectlist_raw, token=ET, ct=projectlist_ct,
    )
    _, imported_projectlist_rows = req(
        "GET", "/api/po?" + urlencode({
            "month": "2026-07",
            "project": projectlist_prefix,
        }),
    )
    imported_projectlist = {
        row["project"]: row for row in imported_projectlist_rows
    }
    expected_projectlist = {
        projectlist_projects["translation"]: {
            "pricing_mode": "per_1000", "word_count": 1200,
            "rate": 200, "amount": 240,
        },
        projectlist_projects["review"]: {
            "pricing_mode": "per_1000", "word_count": 2000,
            "rate": 100, "amount": 200,
        },
        projectlist_projects["mtpe"]: {
            "pricing_mode": "per_1000", "word_count": 1500,
            "rate": 80, "amount": 120,
        },
        projectlist_projects["lqa"]: {
            "pricing_mode": "per_hour", "word_count": 3,
            "rate": 15, "amount": 45,
        },
        projectlist_projects["lqe"]: {
            "pricing_mode": "per_hour", "word_count": 2,
            "rate": 20, "amount": 40,
        },
        projectlist_projects["fixed"]: {
            "pricing_mode": "fixed", "amount": 680,
        },
    }
    chk(
        "Projectlist正式导入按WWC及计价模式写入PO",
        projectlist_import.get("source_format") == "projectlist"
        and projectlist_import.get("preview") is False
        and projectlist_import.get("imported") == 6
        and projectlist_import.get("ready") == 6
        and projectlist_import.get("skipped_settled") == 1
        and projectlist_import.get("skipped_dup_po") == 0
        and len(projectlist_import.get("invalid_rows", [])) == 2
        and set(imported_projectlist) == set(expected_projectlist)
        and all(
            all(
                row.get(field) == value
                for field, value in expected_projectlist[project].items()
                if field != "amount"
            )
            and abs(
                float(row.get("amount", 0))
                - expected_projectlist[project]["amount"]
            ) <= 0.02
            and row.get("amount_ok") is True
            and row.get("source_key")
            for project, row in imported_projectlist.items()
        ),
        {"result": projectlist_import, "rows": imported_projectlist_rows},
    )
    _, repeated_projectlist_import = req(
        "POST", "/api/import/po?projectlist_po_state=all",
        raw=projectlist_raw, token=ET, ct=projectlist_ct,
    )
    repeated_projectlist_rows = req(
        "GET", "/api/po?" + urlencode({
            "month": "2026-07",
            "project": projectlist_prefix,
        }),
    )[1]
    chk(
        "Projectlist同一来源重复导入跳过且不重复写入",
        repeated_projectlist_import.get("imported") == 0
        and repeated_projectlist_import.get("ready") == 0
        and repeated_projectlist_import.get("skipped_settled") == 1
        and repeated_projectlist_import.get("skipped_dup_po") == 6
        and len(repeated_projectlist_import.get("invalid_rows", [])) == 2
        and len(repeated_projectlist_rows) == 6,
        repeated_projectlist_import,
    )
    projectlist_identity_project = "Projectlist稳定身份" + run_tag
    projectlist_identity_raw, projectlist_identity_ct = mp(mk_xlsx_sheets({
        "Projectlist 稳定身份": [
            projectlist_secure_headers,
            [
                projectlist_identity_project, "2026-08-31", "英语",
                filter_translator["name"], "翻译", 0.2, "CNY",
                1000, 200, "8月", "简体中文", "□", "否",
            ],
            [
                projectlist_identity_project, "2026-08-31", "英语",
                projectlist_alias, "翻译", 0.2, "CNY",
                1000, 200, "8月", "简体中文", "□", "否",
            ],
        ],
    }))
    _, projectlist_identity_import = req(
        "POST", "/api/import/po?projectlist_po_state=all",
        raw=projectlist_identity_raw, token=ET, ct=projectlist_identity_ct,
    )
    projectlist_identity_rows = req(
        "GET", "/api/po?" + urlencode({
            "month": "2026-08",
            "project": projectlist_identity_project,
        }),
    )[1]
    identity_source_keys = {
        row.get("source_key") for row in projectlist_identity_rows
    }
    chk(
        "Projectlist正式名与别名同身份多行使用稳定occurrence",
        projectlist_identity_import.get("imported") == 2
        and projectlist_identity_import.get("ready") == 2
        and projectlist_identity_import.get("skipped_dup_po") == 0
        and projectlist_identity_import.get("source_conflicts") == 0
        and not projectlist_identity_import.get("invalid_rows")
        and len(projectlist_identity_rows) == 2
        and len(identity_source_keys) == 2
        and None not in identity_source_keys
        and all(
            row.get("translator_id") == filter_tid
            for row in projectlist_identity_rows
        ),
        {
            "result": projectlist_identity_import,
            "rows": projectlist_identity_rows,
        },
    )
    _, repeated_projectlist_identity = req(
        "POST", "/api/import/po?projectlist_po_state=all",
        raw=projectlist_identity_raw, token=ET, ct=projectlist_identity_ct,
    )
    repeated_projectlist_identity_rows = req(
        "GET", "/api/po?" + urlencode({
            "month": "2026-08",
            "project": projectlist_identity_project,
        }),
    )[1]
    chk(
        "Projectlist同身份多行重传逐行跳过",
        repeated_projectlist_identity.get("imported") == 0
        and repeated_projectlist_identity.get("ready") == 0
        and repeated_projectlist_identity.get("skipped_dup_po") == 2
        and repeated_projectlist_identity.get("source_conflicts") == 0
        and not repeated_projectlist_identity.get("invalid_rows")
        and len(repeated_projectlist_identity_rows) == 2
        and {
            row.get("source_key")
            for row in repeated_projectlist_identity_rows
        } == identity_source_keys,
        repeated_projectlist_identity,
    )
    identity_rows_before_revision = sorted(
        (
            row.get("id"), row.get("source_key"), row.get("word_count"),
            row.get("rate"), row.get("currency"), row.get("amount"),
        )
        for row in repeated_projectlist_identity_rows
    )
    projectlist_revision_raw, projectlist_revision_ct = mp(mk_xlsx_sheets({
        "Projectlist 内容修订": [
            projectlist_secure_headers,
            [
                projectlist_identity_project, "2026-08-31", "英语",
                filter_translator["name"], "翻译", 0.25, "USD",
                1200, None, "8月", "简体中文", "□", "否",
            ],
        ],
    }))
    _, projectlist_revision_import = req(
        "POST", "/api/import/po?projectlist_po_state=all",
        raw=projectlist_revision_raw, token=ET, ct=projectlist_revision_ct,
    )
    projectlist_identity_after_revision = req(
        "GET", "/api/po?" + urlencode({
            "month": "2026-08",
            "project": projectlist_identity_project,
        }),
    )[1]
    identity_rows_after_revision = sorted(
        (
            row.get("id"), row.get("source_key"), row.get("word_count"),
            row.get("rate"), row.get("currency"), row.get("amount"),
        )
        for row in projectlist_identity_after_revision
    )
    revision_conflicts = [
        row
        for row in projectlist_revision_import.get("invalid_rows", [])
        if row.get("source_conflict") is True
    ]
    chk(
        "Projectlist已导入身份内容修订拒绝新增",
        projectlist_revision_import.get("imported") == 0
        and projectlist_revision_import.get("ready") == 0
        and projectlist_revision_import.get("skipped_dup_po") == 0
        and projectlist_revision_import.get("source_conflicts") == 1
        and len(revision_conflicts) == 1
        and all(
            label in revision_conflicts[0].get("error", "")
            for label in ("数量/小时", "费率", "币种")
        )
        and identity_rows_after_revision == identity_rows_before_revision,
        {
            "result": projectlist_revision_import,
            "rows": projectlist_identity_after_revision,
        },
    )
    projectlist_edge_prefix = "Projectlist边界" + run_tag
    projectlist_edge_projects = {
        "rollover": projectlist_edge_prefix + "-跨年",
        "unknown_po": projectlist_edge_prefix + "-未知结算PO",
        "unknown_paid": projectlist_edge_prefix + "-未知已打款",
        "paid_missing_translator": projectlist_edge_prefix + "-已打款无译员",
        "fuzzy_name": projectlist_edge_prefix + "-禁止包含匹配",
        "ambiguous_month": projectlist_edge_prefix + "-超两月月份",
    }
    projectlist_edge_raw, projectlist_edge_ct = mp(mk_xlsx_sheets({
        "Projectlist 边界": [
            ["说明"],
            [
                "项目名称（稿件/LQA批次）", "DDL", "目标语言", "指定译员",
                "工作类型", "翻译费率", "币种", "REPNEW 实际",
                "译员WWC字数", "REPNEW & 小时数", "稿费金额（CNY）",
                "译员PO时间（X月）", "原语言", "结算PO", "已打款",
            ],
            [
                projectlist_edge_projects["rollover"], "2026-12-31", "英语",
                filter_translator["name"], "翻译", 0.2, "CNY", 9900,
                1000, 77, 200, "1月", "简体中文", "□", "否",
            ],
            [
                projectlist_edge_projects["unknown_po"], "2026-12-31", "英语",
                filter_translator["name"], "翻译", 0.2, "CNY", 8800,
                1000, 66, 200, "1月", "简体中文", "待处理", "否",
            ],
            [
                projectlist_edge_projects["unknown_paid"], "2026-12-31", "英语",
                filter_translator["name"], "翻译", 0.2, "CNY", 7700,
                1000, 55, 200, "1月", "简体中文", "□", "待确认",
            ],
            [
                projectlist_edge_projects["paid_missing_translator"],
                "2026-12-31", "英语", "不存在译员" + run_tag,
                "翻译", 0.2, "CNY", 6600, 1000, 44, 200,
                "1月", "简体中文", "□", "是",
            ],
            [
                projectlist_edge_projects["fuzzy_name"], "2026-12-31", "英语",
                filter_translator["name"][1:], "翻译", 0.2, "CNY", 5500,
                1000, 33, 200, "1月", "简体中文", "□", "否",
            ],
            [
                projectlist_edge_projects["ambiguous_month"],
                "2026-03-31", "英语", filter_translator["name"],
                "翻译", 0.2, "CNY", 4400, 1000, 22, 200,
                "6月", "简体中文", "□", "否",
            ],
        ],
    }))
    _, projectlist_edge_preview = req(
        "POST", "/api/import/po?preview=true&projectlist_po_state=all",
        raw=projectlist_edge_raw, token=ET, ct=projectlist_edge_ct,
    )
    edge_preview_by_project = {
        row["project"]: row
        for row in projectlist_edge_preview.get("preview_rows", [])
    }
    edge_preview_errors = {
        row.get("row"): row.get("error")
        for row in projectlist_edge_preview.get("invalid_rows", [])
    }
    chk(
        "Projectlist跨年月份、状态和精确姓名边界预览",
        projectlist_edge_preview.get("ready") == 1
        and projectlist_edge_preview.get("checked_rows") == 0
        and projectlist_edge_preview.get("unchecked_rows") == 4
        and projectlist_edge_preview.get("paid_rows") == 1
        and projectlist_edge_preview.get("settled_rows") == 1
        and projectlist_edge_preview.get("unknown_state_rows") == 2
        and projectlist_edge_preview.get("skipped_settled") == 1
        and set(edge_preview_errors) == {4, 5, 7, 8}
        and "状态无法识别" in edge_preview_errors[4]
        and "状态无法识别" in edge_preview_errors[5]
        and "译员不存在" in edge_preview_errors[7]
        and "缺有效译员PO时间" in edge_preview_errors[8]
        and edge_preview_by_project[
            projectlist_edge_projects["rollover"]
        ].get("settlement_month") == "2027-01"
        and edge_preview_by_project[
            projectlist_edge_projects["rollover"]
        ].get("action") == "import"
        and edge_preview_by_project[
            projectlist_edge_projects["paid_missing_translator"]
        ].get("action") == "skip_settled"
        and edge_preview_by_project[
            projectlist_edge_projects["paid_missing_translator"]
        ].get("translator_id") is None
        and edge_preview_by_project[
            projectlist_edge_projects["paid_missing_translator"]
        ].get("already_paid") is True,
        projectlist_edge_preview,
    )
    _, projectlist_edge_import = req(
        "POST", "/api/import/po?projectlist_po_state=all",
        raw=projectlist_edge_raw, token=ET, ct=projectlist_edge_ct,
    )
    imported_projectlist_edge_rows = req(
        "GET", "/api/po?" + urlencode({
            "month": "2027-01",
            "project": projectlist_edge_prefix,
        }),
    )[1]
    chk(
        "Projectlist边界行正式导入仅写入可确定PO",
        projectlist_edge_import.get("imported") == 1
        and projectlist_edge_import.get("ready") == 1
        and projectlist_edge_import.get("skipped_settled") == 1
        and projectlist_edge_import.get("skipped_dup_po") == 0
        and len(projectlist_edge_import.get("invalid_rows", [])) == 4
        and len(imported_projectlist_edge_rows) == 1
        and imported_projectlist_edge_rows[0].get("project")
        == projectlist_edge_projects["rollover"]
        and imported_projectlist_edge_rows[0].get("settlement_month")
        == "2027-01"
        and abs(
            float(imported_projectlist_edge_rows[0].get("amount", 0)) - 200
        ) <= 0.02,
        {
            "result": projectlist_edge_import,
            "rows": imported_projectlist_edge_rows,
        },
    )

    def po_log_sheet_records(workbook, sheet_name):
        rows = list(workbook[sheet_name].iter_rows(values_only=True))
        headers = tuple(str(value or "").strip() for value in rows[0])
        records = [
            dict(zip(headers, row))
            for row in rows[1:]
            if any(value is not None for value in row)
        ]
        return headers, records

    po_log_batch_fields = {
        "id": ("id", "batch_id", "批次ID"),
        "source_format": ("source_format", "来源格式"),
        "filename": ("filename", "file_name", "文件名"),
        "file_hash": ("file_hash", "文件哈希"),
        "imported": ("imported", "imported_count", "导入数"),
        "duplicate": (
            "skipped_dup_po", "duplicate_count", "重复数",
        ),
        "settled": (
            "skipped_settled", "settled_count", "历史跳过数",
        ),
        "conflict": (
            "source_conflicts", "conflict_count", "来源冲突数",
        ),
        "invalid": ("invalid_count", "invalid", "错误数"),
    }
    po_log_detail_fields = {
        "batch_id": ("batch_id", "批次ID"),
        "source_row": ("source_row", "row", "源行号"),
        "action": ("action", "result", "outcome", "结果"),
        "project": ("project", "项目"),
        "error": ("error", "error_message", "错误"),
        "po_id": ("po_id", "PO ID"),
        "source_key": ("source_key", "来源Key"),
    }

    def po_log_value(record, field_map, field):
        for name in field_map[field]:
            if name in record:
                return record[name]
        return None

    po_log_baseline_bytes = req(
        "GET", "/api/export/po-log", token=ET,
    )[1]
    po_log_baseline_workbook = load_workbook(
        io.BytesIO(po_log_baseline_bytes), data_only=True,
    )
    chk(
        "PO导入日志导出为含两表的有效XLSX",
        {"导入批次", "行级明细"}.issubset(
            po_log_baseline_workbook.sheetnames
        ),
        po_log_baseline_workbook.sheetnames,
    )
    baseline_batch_headers, baseline_po_log_batches = po_log_sheet_records(
        po_log_baseline_workbook, "导入批次",
    )
    baseline_detail_headers, baseline_po_log_details = po_log_sheet_records(
        po_log_baseline_workbook, "行级明细",
    )
    chk(
        "PO导入日志导出包含批次和行级核心字段",
        all(
            any(name in baseline_batch_headers for name in aliases)
            for aliases in po_log_batch_fields.values()
        )
        and all(
            any(name in baseline_detail_headers for name in aliases)
            for aliases in po_log_detail_fields.values()
        ),
        {
            "batch_headers": baseline_batch_headers,
            "detail_headers": baseline_detail_headers,
        },
    )

    po_log_standard_project = "PO日志标准" + run_tag
    po_log_standard_number = "PO-LOG-STANDARD-" + run_tag
    po_log_standard_raw, po_log_standard_ct = mp(mk_xlsx([
        ["译员", "结算月", "项目", "源语言", "目标语言", "角色",
         "字数", "单价", "币种", "状态", "PO号"],
        ["张明", "2027-02", po_log_standard_project, "ZH", "EN", "翻译",
         1000, 180, "CNY", "未开票", po_log_standard_number],
        ["张明", "2027-02", po_log_standard_project + "-重复",
         "ZH", "EN", "翻译", 1000, 180, "CNY", "未开票",
         po_log_standard_number],
        ["不存在译员" + run_tag, "2027-02", po_log_standard_project + "-错误",
         "ZH", "EN", "翻译", 1000, 180, "CNY", "未开票",
         "PO-LOG-BAD-" + run_tag],
    ]))
    po_log_projectlist_new = "Projectlist日志新增" + run_tag
    po_log_projectlist_settled = "Projectlist日志已结算" + run_tag
    po_log_projectlist_invalid = "Projectlist日志错误" + run_tag
    po_log_projectlist_raw, po_log_projectlist_ct = mp(mk_xlsx_sheets({
        "Projectlist 日志": [
            projectlist_secure_headers,
            [
                projectlist_identity_project, "2026-08-31", "英语",
                filter_translator["name"], "翻译", 0.2, "CNY",
                1000, 200, "8月", "简体中文", "□", "否",
            ],
            [
                projectlist_identity_project, "2026-08-31", "英语",
                projectlist_alias, "翻译", 0.25, "USD",
                1200, None, "8月", "简体中文", "□", "否",
            ],
            [
                po_log_projectlist_new, "2026-08-31", "英语",
                filter_translator["name"], "翻译", 0.2, "CNY",
                1000, 200, "8月", "简体中文", "□", "否",
            ],
            [
                po_log_projectlist_settled, "2026-08-31", "英语",
                "日志不存在译员" + run_tag, "翻译", 0.2, "CNY",
                1000, 200, "8月", "简体中文", "✅", "否",
            ],
            [
                po_log_projectlist_invalid, "2026-08-31", "英语",
                filter_translator["name"], "配音", 50, "CNY",
                1000, 50, "8月", "简体中文", "□", "否",
            ],
        ],
    }))
    req(
        "POST", "/api/import/po?preview=true",
        raw=po_log_standard_raw, token=ET, ct=po_log_standard_ct,
    )
    req(
        "POST", "/api/import/po?preview=true&projectlist_po_state=all",
        raw=po_log_projectlist_raw, token=ET, ct=po_log_projectlist_ct,
    )
    po_log_after_preview_workbook = load_workbook(io.BytesIO(req(
        "GET", "/api/export/po-log", token=ET,
    )[1]), data_only=True)
    after_preview_batches = po_log_sheet_records(
        po_log_after_preview_workbook, "导入批次",
    )[1]
    after_preview_details = po_log_sheet_records(
        po_log_after_preview_workbook, "行级明细",
    )[1]
    chk(
        "PO预览不产生批次或行级日志",
        len(after_preview_batches) == len(baseline_po_log_batches)
        and len(after_preview_details) == len(baseline_po_log_details),
        {
            "before": (
                len(baseline_po_log_batches), len(baseline_po_log_details),
            ),
            "after": (len(after_preview_batches), len(after_preview_details)),
        },
    )

    _, po_log_standard_import = req(
        "POST", "/api/import/po",
        raw=po_log_standard_raw, token=ET, ct=po_log_standard_ct,
    )
    _, po_log_projectlist_import = req(
        "POST", "/api/import/po?projectlist_po_state=all",
        raw=po_log_projectlist_raw, token=ET, ct=po_log_projectlist_ct,
    )
    final_po_log_workbook = load_workbook(io.BytesIO(req(
        "GET", "/api/export/po-log", token=ET,
    )[1]), data_only=True)
    _, final_po_log_batches = po_log_sheet_records(
        final_po_log_workbook, "导入批次",
    )
    _, final_po_log_details = po_log_sheet_records(
        final_po_log_workbook, "行级明细",
    )
    standard_log_batches = [
        row for row in final_po_log_batches
        if po_log_value(
            row, po_log_batch_fields, "file_hash",
        ) == po_log_standard_import.get("file_hash")
    ]
    projectlist_log_batches = [
        row for row in final_po_log_batches
        if po_log_value(
            row, po_log_batch_fields, "file_hash",
        ) == po_log_projectlist_import.get("file_hash")
    ]
    standard_log_batch_id = po_log_value(
        standard_log_batches[0], po_log_batch_fields, "id",
    ) if len(standard_log_batches) == 1 else None
    projectlist_log_batch_id = po_log_value(
        projectlist_log_batches[0], po_log_batch_fields, "id",
    ) if len(projectlist_log_batches) == 1 else None
    standard_log_details = [
        row for row in final_po_log_details
        if po_log_value(row, po_log_detail_fields, "batch_id")
        == standard_log_batch_id
    ]
    projectlist_log_details = [
        row for row in final_po_log_details
        if po_log_value(row, po_log_detail_fields, "batch_id")
        == projectlist_log_batch_id
    ]

    def po_log_outcomes(rows):
        return {
            int(po_log_value(row, po_log_detail_fields, "source_row")):
            po_log_value(row, po_log_detail_fields, "action")
            for row in rows
        }

    standard_log_outcomes = po_log_outcomes(standard_log_details)
    projectlist_log_outcomes = po_log_outcomes(projectlist_log_details)
    chk(
        "标准PO和Projectlist正式导入各产生一个批次日志",
        len(final_po_log_batches) == len(baseline_po_log_batches) + 2
        and len(final_po_log_details) == len(baseline_po_log_details) + 8
        and len(standard_log_batches) == 1
        and len(projectlist_log_batches) == 1
        and po_log_value(
            standard_log_batches[0], po_log_batch_fields, "source_format",
        ) == "standard"
        and po_log_value(
            projectlist_log_batches[0], po_log_batch_fields, "source_format",
        ) == "projectlist"
        and po_log_value(
            standard_log_batches[0], po_log_batch_fields, "filename",
        ) == "x.xlsx"
        and po_log_value(
            projectlist_log_batches[0], po_log_batch_fields, "filename",
        ) == "x.xlsx",
        {
            "standard": standard_log_batches,
            "projectlist": projectlist_log_batches,
        },
    )
    chk(
        "PO导入批次统计与单一行级结果一致",
        po_log_standard_import.get("imported") == 1
        and po_log_standard_import.get("skipped_dup_po") == 1
        and len(po_log_standard_import.get("invalid_rows", [])) == 1
        and po_log_projectlist_import.get("imported") == 1
        and po_log_projectlist_import.get("skipped_dup_po") == 1
        and po_log_projectlist_import.get("skipped_settled") == 1
        and po_log_projectlist_import.get("source_conflicts") == 1
        and len(po_log_projectlist_import.get("invalid_rows", [])) == 2
        and standard_log_outcomes == {
            2: "imported", 3: "skip_duplicate", 4: "invalid",
        }
        and projectlist_log_outcomes == {
            2: "skip_duplicate", 3: "source_conflict",
            4: "imported", 5: "skip_settled", 6: "invalid",
        }
        and len(standard_log_details) == len(standard_log_outcomes)
        and len(projectlist_log_details) == len(projectlist_log_outcomes)
        and all(
            int(po_log_value(batch, po_log_batch_fields, field)) == expected
            for batch, expected_values in (
                (standard_log_batches[0], (1, 1, 0, 0, 1)),
                (projectlist_log_batches[0], (1, 1, 1, 1, 1)),
            )
            for field, expected in zip(
                ("imported", "duplicate", "settled", "conflict", "invalid"),
                expected_values,
            )
        ),
        {
            "standard_result": po_log_standard_import,
            "projectlist_result": po_log_projectlist_import,
            "standard_rows": standard_log_details,
            "projectlist_rows": projectlist_log_details,
        },
    )
    standard_log_by_row = {
        int(po_log_value(row, po_log_detail_fields, "source_row")): row
        for row in standard_log_details
    }
    projectlist_log_by_row = {
        int(po_log_value(row, po_log_detail_fields, "source_row")): row
        for row in projectlist_log_details
    }
    chk(
        "PO日志XLSX含项目、错误及来源代表性数据",
        po_log_value(
            standard_log_by_row[2], po_log_detail_fields, "project",
        ) == po_log_standard_project
        and "译员不存在" in str(po_log_value(
            standard_log_by_row[4], po_log_detail_fields, "error",
        ))
        and po_log_value(
            projectlist_log_by_row[4], po_log_detail_fields, "project",
        ) == po_log_projectlist_new
        and "发生变化" in str(po_log_value(
            projectlist_log_by_row[3], po_log_detail_fields, "error",
        ))
        and "工作类型" in str(po_log_value(
            projectlist_log_by_row[6], po_log_detail_fields, "error",
        ))
        and po_log_value(
            projectlist_log_by_row[2], po_log_detail_fields, "source_key",
        ),
        {
            "standard": standard_log_details,
            "projectlist": projectlist_log_details,
        },
    )
    _, cumulative_summary = req("GET", "/api/po/summary?month=2026-11")
    chk(
        "PO汇总同时返回本月和跨月累计未付",
        cumulative_summary["by_currency"]["CNY"]["unpaid"] >= 1200
        and cumulative_summary["cumulative_unpaid_by_currency"]["CNY"]
        >= cumulative_summary["by_currency"]["CNY"]["unpaid"],
        cumulative_summary,
    )

    _, wechat_account = req(
        "POST", f"/api/translators/{filter_tid}/payment-accounts",
        {
            "method": "wechat",
            "account_name": filter_translator["name"],
            "account_number": "wx-" + run_tag,
            "is_default": True,
        },
        token=ET,
    )
    _, usd_account = req(
        "POST", f"/api/translators/{filter_tid}/payment-accounts",
        {
            "method": "corporate_usd",
            "account_name": "Filter Studio",
            "account_number": "12345678",
            "bank_name": "Test Bank",
            "bank_address": "1 Test Road",
            "swift_code": "TESTUS33",
            "routing_code": "110000",
            "is_default": True,
        },
        token=ET,
    )
    _, alipay_account = req(
        "POST", f"/api/translators/{filter_tid}/payment-accounts",
        {
            "method": "alipay",
            "account_name": filter_translator["name"],
            "account_number": "alipay-" + run_tag,
        },
        token=ET,
    )
    _, personal_account = req(
        "POST", f"/api/translators/{filter_tid}/payment-accounts",
        {
            "method": "personal_bank",
            "account_name": filter_translator["name"],
            "account_number": "6222000011112222",
            "bank_name": "测试银行",
        },
        token=ET,
    )
    _, cny_account = req(
        "POST", f"/api/translators/{filter_tid}/payment-accounts",
        {
            "method": "corporate_cny",
            "account_name": "Filter Studio CN",
            "account_number": "6222000099990000",
            "bank_name": "测试银行",
            "tax_id": "91310000TEST",
        },
        token=ET,
    )
    accounts = req(
        "GET", f"/api/translators/{filter_tid}/payment-accounts",
    )[1]
    chk(
        "五种支付方式可并存且暂不设置默认账户、敏感账号脱敏",
        sum(1 for account in accounts if account["is_default"]) == 0
        and {account["method"] for account in accounts}
        == {
            "wechat", "alipay", "personal_bank",
            "corporate_cny", "corporate_usd",
        }
        and next(
            account for account in accounts if account["id"] == usd_account["id"]
        )["account_number"].endswith("5678"),
        accounts,
    )
    _, updated_usd_account = req(
        "PUT",
        f"/api/translators/{filter_tid}/payment-accounts/{usd_account['id']}",
        {
            "method": "corporate_usd",
            "account_name": "Filter Studio",
            "bank_name": "Test Bank",
            "bank_address": "2 Test Road",
            "swift_code": "TESTUS33",
            "routing_code": "220000",
            "is_default": True,
        },
        token=ET,
    )
    chk(
        "支付账户可编辑且空账号字段保留原加密值",
        updated_usd_account.get("bank_address") == "2 Test Road"
        and updated_usd_account.get("account_number", "").endswith("5678"),
        updated_usd_account,
    )
    revealed_account = req(
        "GET",
        f"/api/translators/{filter_tid}/payment-accounts/{usd_account['id']}/reveal",
        token=ET,
    )[1]
    chk(
        "支付账户明文仅编辑角色可审计读取",
        revealed_account.get("account_number") == "12345678"
        and code(lambda: req(
            "GET",
            f"/api/translators/{filter_tid}/payment-accounts/{usd_account['id']}/reveal",
            token=BT,
        )) == 403,
        revealed_account,
    )
    _, minimal_payment = req(
        "POST", f"/api/translators/{filter_tid}/payment-accounts",
        {
            "method": "corporate_cny",
            "remarks": "仅备注也可保存",
        },
        token=ET,
    )
    chk(
        "支付方式字段均非必填但至少填写一项信息",
        minimal_payment.get("remarks") == "仅备注也可保存"
        and code(lambda: req(
            "POST", f"/api/translators/{filter_tid}/payment-accounts",
            {"method": "corporate_cny"},
            token=ET,
        )) == 400,
        minimal_payment,
    )
    png_data = b"\x89PNG\r\n\x1a\nacceptance"
    qr_raw, qr_ct = mp_file(png_data, "wechat.png", "image/png")
    _, qr_result = req(
        "POST",
        f"/api/translators/{filter_tid}/payment-accounts/{wechat_account['id']}/qr",
        raw=qr_raw, token=ET, ct=qr_ct,
    )
    qr_download = req(
        "GET",
        f"/api/translators/{filter_tid}/payment-accounts/{wechat_account['id']}/qr",
        token=BT,
    )[1]
    chk(
        "微信收款码可上传并由登录角色下载",
        qr_result.get("has_qr") is True and qr_download == png_data,
        qr_result,
    )
    chk(
        "收款码未登录不可下载",
        code(lambda: req(
            "GET",
            f"/api/translators/{filter_tid}/payment-accounts/{wechat_account['id']}/qr",
        )) == 403,
    )
    req(
        "DELETE",
        f"/api/translators/{filter_tid}/payment-accounts/{usd_account['id']}",
        token=ET,
    )
    accounts_after_delete = req(
        "GET", f"/api/translators/{filter_tid}/payment-accounts",
    )[1]
    chk(
        "删除账户后不自动指定默认账户",
        len(accounts_after_delete) == 5
        and accounts_after_delete[0]["id"] == wechat_account["id"]
        and accounts_after_delete[0]["is_default"] is False
        and sum(
            1 for account in accounts_after_delete if account["is_default"]
        ) == 0,
        accounts_after_delete,
    )

    attachment_data = b"PK\x03\x04qualification"
    attachment_raw, attachment_ct = mp_file(
        attachment_data,
        "sample.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        {"category": "sample"},
    )
    _, uploaded_attachment = req(
        "POST", f"/api/translators/{filter_tid}/attachments",
        raw=attachment_raw, token=ET, ct=attachment_ct,
    )
    attachments = req(
        "GET", f"/api/translators/{filter_tid}/attachments", token=BT,
    )[1]
    attachment_download = req(
        "GET",
        f"/api/translators/{filter_tid}/attachments/{uploaded_attachment['id']}",
        token=BT,
    )[1]
    chk(
        "资质附件保存元数据且登录角色可下载",
        len(attachments) == 1
        and attachments[0]["category"] == "sample"
        and len(attachments[0]["sha256"]) == 64
        and attachment_download == attachment_data,
        attachments,
    )
    chk(
        "资质附件列表未登录403",
        code(lambda: req(
            "GET", f"/api/translators/{filter_tid}/attachments",
        )) == 403,
    )
    fake_pdf_raw, fake_pdf_ct = mp_file(
        b"not-a-pdf", "fake.pdf", "application/pdf", {"category": "certificate"},
    )
    chk(
        "伪造扩展名附件400",
        code(lambda: req(
            "POST", f"/api/translators/{filter_tid}/attachments",
            raw=fake_pdf_raw, token=ET, ct=fake_pdf_ct,
        )) == 400,
    )
    req(
        "DELETE",
        f"/api/translators/{filter_tid}/attachments/{uploaded_attachment['id']}",
        token=ET,
    )
    chk(
        "资质附件删除后元数据和文件均不可访问",
        req(
            "GET", f"/api/translators/{filter_tid}/attachments", token=ET,
        )[1] == []
        and code(lambda: req(
            "GET",
            f"/api/translators/{filter_tid}/attachments/{uploaded_attachment['id']}",
            token=ET,
        )) == 404,
    )

    print()
    print("结果:", f"{sum(ok)}/{len(ok)} 全通过 ✅" if all(ok) else f"{sum(ok)}/{len(ok)}（有失败）")
    rc = 0 if all(ok) else 1
    stop_isolated_server()
    return rc


if __name__ == "__main__":
    sys.exit(main())
