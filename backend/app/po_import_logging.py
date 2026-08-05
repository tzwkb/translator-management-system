"""Persistent PO import outcomes and audit-friendly workbook export."""

from dataclasses import dataclass
from datetime import datetime
import math
import re
from typing import Optional

from sqlalchemy.orm import Session

from .models import POImportBatch, POImportRowLog


ACTION_LABELS = {
    "imported": "已导入",
    "skip_duplicate": "跳过重复",
    "skip_settled": "跳过历史",
    "source_conflict": "来源冲突",
    "invalid": "错误",
    "ignored": "未选中/忽略",
}
ALLOWED_ACTIONS = frozenset(ACTION_LABELS)
_ILLEGAL_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]")


def _clean_text(value, max_length=None):
    if value is None:
        return None
    cleaned = _ILLEGAL_CONTROL_CHARS.sub("", str(value))
    return cleaned[:max_length] if max_length else cleaned


def safe_import_filename(value):
    name = str(value or "po.xlsx").replace("\\", "/").rsplit("/", 1)[-1]
    return _clean_text(name.strip(), 255) or "po.xlsx"


@dataclass
class POImportOutcome:
    action: str
    sheet: str
    source_row: int
    selected: bool = True
    po_id: Optional[int] = None
    translator_id: Optional[int] = None
    translator_name: Optional[str] = None
    project: Optional[str] = None
    settlement_month: Optional[str] = None
    role: Optional[str] = None
    source_lang: Optional[str] = None
    target_lang: Optional[str] = None
    pricing_mode: Optional[str] = None
    word_count: Optional[float] = None
    rate: Optional[float] = None
    amount: Optional[float] = None
    source_fee_cny: Optional[float] = None
    currency: Optional[str] = None
    source_key: Optional[str] = None
    error_code: Optional[str] = None
    error_message: Optional[str] = None

    def model_kwargs(self, batch_id):
        return {
            "batch_id": batch_id,
            "sheet": _clean_text(self.sheet, 255),
            "source_row": self.source_row,
            "action": self.action,
            "selected": self.selected,
            "po_id": self.po_id,
            "translator_id": self.translator_id,
            "translator_name": _clean_text(self.translator_name, 200),
            "project": _clean_text(self.project, 200),
            "settlement_month": _clean_text(self.settlement_month, 7),
            "role": _clean_text(self.role, 50),
            "source_lang": _clean_text(self.source_lang, 20),
            "target_lang": _clean_text(self.target_lang, 20),
            "pricing_mode": _clean_text(self.pricing_mode, 20),
            "word_count": _finite_number(self.word_count),
            "rate": _finite_number(self.rate),
            "amount": _finite_number(self.amount),
            "source_fee_cny": _finite_number(self.source_fee_cny),
            "currency": _clean_text(self.currency, 10),
            "source_key": _clean_text(self.source_key, 64),
            "error_code": _clean_text(self.error_code, 50),
            "error_message": _clean_text(self.error_message, 4000),
        }


def persist_po_import_log(
    session: Session,
    *,
    created_by,
    source_format,
    parser_version,
    file_name,
    file_hash,
    sheet=None,
    header_row=None,
    projectlist_po_state=None,
    outcomes=(),
):
    outcomes = list(outcomes)
    row_keys = set()
    action_counts = {action: 0 for action in ALLOWED_ACTIONS}
    for outcome in outcomes:
        if outcome.action not in ALLOWED_ACTIONS:
            raise ValueError(f"非法 PO 导入日志 action：{outcome.action}")
        if not outcome.sheet or outcome.source_row < 1:
            raise ValueError("PO 导入日志必须包含工作表和正源行号")
        row_key = (outcome.sheet, outcome.source_row)
        if row_key in row_keys:
            raise ValueError(f"PO 导入日志源行重复：{row_key}")
        row_keys.add(row_key)
        if outcome.action == "ignored" and outcome.selected:
            raise ValueError("ignored 行必须标记为未选中")
        if outcome.action != "ignored" and not outcome.selected:
            raise ValueError("未选中行只能记录为 ignored")
        action_counts[outcome.action] += 1

    batch = POImportBatch(
        created_by=_clean_text(created_by, 50),
        source_format=_clean_text(source_format, 20),
        parser_version=_clean_text(parser_version, 30),
        file_name=safe_import_filename(file_name),
        file_hash=_clean_text(file_hash, 64),
        sheet=_clean_text(sheet, 255),
        header_row=header_row,
        projectlist_po_state=_clean_text(projectlist_po_state, 20),
        selected_rows=len(outcomes) - action_counts["ignored"],
        ignored_rows=action_counts["ignored"],
        imported=action_counts["imported"],
        skipped_duplicate=action_counts["skip_duplicate"],
        skipped_settled=action_counts["skip_settled"],
        source_conflicts=action_counts["source_conflict"],
        invalid_count=action_counts["invalid"],
        created_at=datetime.now(),
    )
    session.add(batch)
    session.flush()
    session.add_all([
        POImportRowLog(**outcome.model_kwargs(batch.id))
        for outcome in outcomes
    ])
    return batch


def _excel_text(value):
    if value is None:
        return None
    text = _clean_text(value)
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _finite_number(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError, OverflowError):
        return None
    return number if math.isfinite(number) else None


def _number(value):
    return _finite_number(value)


def build_po_log_workbook(batches, row_logs):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    batches = list(batches)
    row_logs = list(row_logs)
    workbook = Workbook()
    batch_sheet = workbook.active
    batch_sheet.title = "导入批次"
    detail_sheet = workbook.create_sheet("行级明细")

    batch_headers = [
        "批次ID", "导入时间", "操作者", "来源格式", "解析器版本",
        "文件名", "文件哈希", "工作表", "表头行", "Projectlist筛选",
        "选中行", "忽略/未选中行", "导入数", "重复数", "历史跳过数",
        "来源冲突数", "错误数", "状态",
    ]
    batch_sheet.append(batch_headers)
    for batch in batches:
        batch_sheet.append([
            batch.id,
            batch.created_at,
            _excel_text(batch.created_by),
            _excel_text(batch.source_format),
            _excel_text(batch.parser_version),
            _excel_text(batch.file_name),
            _excel_text(batch.file_hash),
            _excel_text(batch.sheet),
            batch.header_row,
            _excel_text(batch.projectlist_po_state),
            batch.selected_rows,
            batch.ignored_rows,
            batch.imported,
            batch.skipped_duplicate,
            batch.skipped_settled,
            batch.source_conflicts,
            batch.invalid_count,
            "完成但有问题"
            if batch.source_conflicts or batch.invalid_count
            else "完成",
        ])

    batch_map = {batch.id: batch for batch in batches}
    detail_headers = [
        "日志ID", "批次ID", "导入时间", "文件名", "工作表", "源行号",
        "action", "结果", "是否选中", "PO ID", "译员ID", "译员姓名", "项目",
        "结算月", "工作类型", "原语言", "目标语言", "计价方式",
        "数量/小时", "费率", "最终金额", "源CNY稿费", "币种", "来源Key",
        "错误代码", "错误",
    ]
    detail_sheet.append(detail_headers)
    for row in row_logs:
        if row.batch_id not in batch_map:
            raise ValueError(f"PO 行日志缺导入批次：{row.batch_id}")
        batch = batch_map[row.batch_id]
        detail_sheet.append([
            row.id,
            row.batch_id,
            batch.created_at,
            _excel_text(batch.file_name),
            _excel_text(row.sheet),
            row.source_row,
            row.action,
            ACTION_LABELS.get(row.action, row.action),
            "是" if row.selected else "否",
            row.po_id,
            row.translator_id,
            _excel_text(row.translator_name),
            _excel_text(row.project),
            _excel_text(row.settlement_month),
            _excel_text(row.role),
            _excel_text(row.source_lang),
            _excel_text(row.target_lang),
            _excel_text(row.pricing_mode),
            _number(row.word_count),
            _number(row.rate),
            _number(row.amount),
            _number(row.source_fee_cny),
            _excel_text(row.currency),
            _excel_text(row.source_key),
            _excel_text(row.error_code),
            _excel_text(row.error_message),
        ])

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    header_border = Border(bottom=Side(style="medium", color="163A5C"))
    status_fills = {
        "已导入": "E2F0D9",
        "跳过重复": "FFF2CC",
        "跳过历史": "D9EAF7",
        "来源冲突": "FCE4D6",
        "错误": "F4CCCC",
        "完成": "E2F0D9",
        "完成但有问题": "FFF2CC",
        "失败": "F4CCCC",
    }

    for sheet, widths in (
        (batch_sheet, [10, 20, 14, 14, 16, 30, 68, 28, 10, 18,
                       10, 14, 10, 12, 12, 12, 10, 18]),
        (detail_sheet, [10, 10, 20, 30, 28, 10, 18, 16, 10, 10, 10,
                        20, 28, 12, 14, 12, 12, 14, 14, 14, 14, 14,
                        10, 68, 18, 48]),
    ):
        sheet.sheet_view.showGridLines = False
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.border = header_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
        sheet.row_dimensions[1].height = 24
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[
                sheet.cell(row=1, column=column).column_letter
            ].width = width

    for cell in batch_sheet["B"][1:]:
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    for row in batch_sheet.iter_rows(min_row=2, min_col=11, max_col=17):
        for cell in row:
            cell.number_format = "#,##0"
    for cell in batch_sheet["R"][1:]:
        if cell.value in status_fills:
            cell.fill = PatternFill("solid", fgColor=status_fills[cell.value])

    for cell in detail_sheet["C"][1:]:
        cell.number_format = "yyyy-mm-dd hh:mm:ss"
    for row in detail_sheet.iter_rows(min_row=2):
        row[18].number_format = "#,##0.00"
        row[19].number_format = "#,##0.000000"
        row[20].number_format = "#,##0.00"
        row[21].number_format = "#,##0.00"
        result_cell = row[7]
        if result_cell.value in status_fills:
            result_cell.fill = PatternFill(
                "solid", fgColor=status_fills[result_cell.value]
            )
        row[25].alignment = Alignment(vertical="top", wrap_text=True)

    return workbook
